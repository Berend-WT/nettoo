#!/usr/bin/env python3
"""Bouwt het beoordelingswerkblad met de gevonden fotovoorstellen erin.

Toont per daily de vraag naast de voorgestelde foto, zodat een redacteur in één
oogopslag ziet of het beeld ergens op slaat. Alleen de kolom Akkoord hoeft
ingevuld te worden.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
VOORSTELLEN = ROOT / "fotos" / "daily_fotovoorstellen.json"
OUTPUT = ROOT / "fotos" / "daily_fotos_beoordeling.xlsx"

FONT = "Arial"
KOP_FILL = PatternFill("solid", fgColor="14163B")
INVUL_FILL = PatternFill("solid", fgColor="FFFFCC")
FOUT_FILL = PatternFill("solid", fgColor="FDE8E8")

KOLOMMEN = [
    ("Daily", 7), ("Datum", 11), ("Vraag", 46), ("Antw.", 9),
    ("Zoekterm", 18), ("Foto", 24), ("Titel op Commons", 34),
    ("Maker", 22), ("Licentie", 14), ("Akkoord", 10), ("Opmerking", 26),
]


def main() -> None:
    rows = json.loads(VOORSTELLEN.read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Beoordeling"

    ws.append([naam for naam, _ in KOLOMMEN])
    for i, (_, breedte) in enumerate(KOLOMMEN, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breedte
        cel = ws.cell(1, i)
        cel.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cel.fill = KOP_FILL
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"

    for r in rows:
        gelukt = r.get("status") == "ok"
        ws.append([
            r.get("nummer"), r.get("datum"), r.get("vraag"), r.get("antwoord"),
            r.get("zoekterm"), "", r.get("titel", ""), r.get("maker", ""),
            r.get("licentie", ""), "", "" if gelukt else r.get("status", ""),
        ])
        rij = ws.max_row
        ws.row_dimensions[rij].height = 92
        for i in range(1, len(KOLOMMEN) + 1):
            cel = ws.cell(rij, i)
            cel.font = Font(name=FONT, size=10)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
            if i == 10:
                cel.fill = INVUL_FILL
            elif not gelukt:
                cel.fill = FOUT_FILL

        bestand = ROOT / (r.get("bestand") or "")
        if gelukt and bestand.is_file():
            try:
                plaatje = XLImage(str(bestand))
                # Vaste hoogte, breedte meeschalen: anders wordt de rij een rommel.
                verhouding = plaatje.width / plaatje.height if plaatje.height else 1
                plaatje.height = 88
                plaatje.width = int(88 * verhouding)
                ws.add_image(plaatje, f"F{rij}")
            except Exception as err:
                ws.cell(rij, 11).value = f"thumbnail mislukt: {err}"

    uitleg = wb.create_sheet("Uitleg")
    for regel in UITLEG:
        uitleg.append([regel])
    uitleg.column_dimensions["A"].width = 100
    for rij in range(1, uitleg.max_row + 1):
        uitleg.cell(rij, 1).font = Font(name=FONT, size=11)
        uitleg.cell(rij, 1).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUTPUT)
    gelukt = sum(1 for r in rows if r.get("status") == "ok")
    print(f"Geschreven: {OUTPUT.relative_to(ROOT)} ({gelukt} met foto, {len(rows) - gelukt} zonder)")


UITLEG = [
    "Fotovoorstellen bij de dailies — beoordeling",
    "",
    "Vul alleen de gele kolom Akkoord: ja of nee.",
    "",
    "De foto's zijn automatisch gezocht op Wikimedia Commons, op basis van een zoekterm die is",
    "afgeleid uit de vraag. Vragen met een eigennaam (Waterloo, Buckingham Palace) gaan bijna altijd",
    "goed; algemene begrippen vaak niet. Reken dus op missers.",
    "",
    "Waar je op let:",
    "De foto mag laten zien WAAR de vraag over gaat, maar niet helpen het GETAL te schatten.",
    "Een piano bij 'hoeveel toetsen heeft een piano' is goed; je telt er geen toetsen op.",
    "Een foto waarop je het antwoord kunt natellen of aflezen is fout.",
    "Slaat de foto nergens op? Zet nee. Een lege daily valt gewoon terug op een sfeerfoto.",
    "",
    "Rood gemarkeerde rijen zijn dailies waarvoor niets gevonden is; daar hoef je niets te doen.",
    "",
    "Licenties zijn al gecontroleerd: alleen CC BY, CC BY-SA, CC0 en publiek domein. Maker en",
    "bronvermelding worden automatisch meegenomen naar de site.",
]


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Netto — schrap de vijf vragen met antwoord nul.

Draaien:  python tools/schrap_nulantwoorden.py [--schrijf]

Deze vijf ontbraken in de eerdere schrapronde omdat ze wel een bruikbare bron
hadden. Antwoord nul maakt elke deling stuk en verandert bij optellen niets, dus
in een rekenspel zijn ze onbruikbaar hoe juist het antwoord ook is.

Vier zijn bovendien inhoudelijk kapot — een ervan bevat letterlijk de correctie
van de schrijver ("sorry dat zit niet in Euro..."). De vijfde is juist perfect
geformuleerd: het vriespunt van water is nul graden Celsius. Die sneuvelt puur
op de rekenkunde.

Zoals eerder: alleen de vragenbank en het reviewblad. De puzzelbestanden blijven
ongemoeid; wat erdoor breekt komt in vragen/puzzels_om_te_herbouwen.xlsx.
"""

import json
import os
import shutil
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK = os.path.join(WORTEL, 'vragen', '1000+ vragen netjes gecategoriseerd.xlsx')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
GESCHRAPT = os.path.join(WORTEL, 'vragen', 'geschrapte_vragen.xlsx')
HERBOUW = os.path.join(WORTEL, 'vragen', 'puzzels_om_te_herbouwen.xlsx')
DATA = os.path.join(WORTEL, 'data')

BESTANDEN = ['netto_frontend_puzzles.js', 'netto_breinkrakers.js',
             'netto_race_sets.js', 'netto_race_pool.js']


def laad(pad):
    t = open(pad, encoding='utf-8').read()
    i = t.index('=', t.index('window.')) + 1
    return json.loads(t[i:].strip().rstrip(';'))


def vragen_in(p):
    uit = []
    for i in (1, 2, 3, 4):
        g = p.get(f'q{i}')
        if isinstance(g, dict) and 'label' in g:
            uit.append(g['label'])
        elif f'q{i}_label' in p:
            uit.append(p[f'q{i}_label'])
    return uit


def alle_puzzels(data):
    if isinstance(data, list):
        return [('', p) for p in data]
    return [(k, p) for k, v in data.items() if isinstance(v, list)
            for p in v if isinstance(p, dict)]


def main():
    schrijf = '--schrijf' in sys.argv
    r = pd.read_excel(REVIEW, sheet_name='Vragen')
    r['_a'] = pd.to_numeric(r['Antwoord'], errors='coerce')
    weg = r[r['_a'] == 0].drop(columns=['_a'])
    labels = {str(v): int(n) for v, n in zip(weg['Vraag NL'], weg['Nr'])}

    print(f'te schrappen: {len(weg)}')
    for _, x in weg.iterrows():
        print(f"  {int(x['Nr']):5d} [{x['In gebruik']}] {str(x['Vraag NL'])[:70]}")

    rijen = []
    for naam in BESTANDEN:
        for sleutel, p in alle_puzzels(laad(os.path.join(DATA, naam))):
            for label in vragen_in(p):
                if label in labels:
                    rijen.append({
                        'Bestand': naam.replace('netto_', '').replace('.js', ''),
                        'Lijst': sleutel, 'Puzzel': p.get('id'),
                        'Datum': p.get('date', ''),
                        'Som': p.get('calculation') or p.get('formula', ''),
                        'Vraag Nr': labels[label], 'Vraag': label,
                        'Oud antwoord': 0, 'Vervangers met dat antwoord': 0,
                        'Advies': 'vraag geschrapt, puzzel laten vervallen'})
    nieuw = pd.DataFrame(rijen)
    print(f'\npuzzels die dit raakt: {len(nieuw)}')
    if len(nieuw):
        print(nieuw.groupby('Bestand').size().to_string())

    if not schrijf:
        print('\n(proefdraai — voeg --schrijf toe om op te slaan)')
        return

    # geschrapte vragen aanvullen in plaats van overschrijven
    if os.path.exists(GESCHRAPT):
        oud = pd.read_excel(GESCHRAPT, sheet_name='Geschrapt')
        pd.concat([oud, weg], ignore_index=True).to_excel(
            GESCHRAPT, index=False, sheet_name='Geschrapt')
    else:
        weg.to_excel(GESCHRAPT, index=False, sheet_name='Geschrapt')

    if len(nieuw) and os.path.exists(HERBOUW):
        oud = pd.read_excel(HERBOUW, sheet_name='Te herbouwen')
        samen = pd.concat([oud, nieuw], ignore_index=True).drop_duplicates(
            subset=['Bestand', 'Lijst', 'Puzzel', 'Vraag Nr'])
        samen.sort_values(['Bestand', 'Lijst', 'Puzzel']).to_excel(
            HERBOUW, index=False, sheet_name='Te herbouwen')

    for pad in (REVIEW, BANK):
        shutil.copy2(pad, os.path.join(
            os.path.dirname(pad),
            f'_backup_nul_{date.today():%Y-%m-%d}_{os.path.basename(pad)}'))
        boek = load_workbook(pad)
        for blad in boek.worksheets:
            kop = [c.value for c in blad[1]]
            if 'Vraag NL' not in kop:
                continue
            i_v = kop.index('Vraag NL') + 1
            teweg = [rij for rij in range(2, blad.max_row + 1)
                     if blad.cell(row=rij, column=i_v).value in labels]
            for rij in sorted(teweg, reverse=True):
                blad.delete_rows(rij)
            if teweg:
                print(f'{os.path.basename(pad)} / {blad.title}: {len(teweg)} weg')
        boek.save(pad)

    na = pd.read_excel(REVIEW, sheet_name='Vragen')
    print(f'\nvragen over: {len(na)}')


if __name__ == '__main__':
    main()

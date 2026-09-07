# -*- coding: utf-8 -*-
"""Netto — zet de vervangende bronnen door naar het reviewblad.

Draaien:  python tools/verwerk_vervangen_bronnen.py

De dode link blijft staan in de kolom "Bron (bestaand)", zodat zichtbaar blijft
wat er ooit stond. De nieuwe bron komt in "Bron (geverifieerd)". Antwoorden
worden niet aangepast: deze vragen zitten in puzzels en een ander getal breekt
de som.
"""

import os
import shutil
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vervangen_bronnen import BRONNEN

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DOEL = os.path.join(WORTEL, 'vragen', 'vervangen_bronnen.xlsx')


def main():
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    rijen = []
    for _, r in d.iterrows():
        nr = int(r['Nr'])
        if nr not in BRONNEN:
            continue
        bron, oordeel, opmerking = BRONNEN[nr]
        rijen.append({
            'Nr': nr, 'Oordeel': oordeel, 'In gebruik': r['In gebruik'],
            'Vraag NL': r['Vraag NL'], 'Antwoord': r['Antwoord'],
            'Dode link (oud)': r['Bron (bestaand)'],
            'Nieuwe bron': bron, 'Toelichting': opmerking, 'Jouw besluit': None,
        })
    uit = pd.DataFrame(rijen).sort_values(['Oordeel', 'In gebruik', 'Nr'])
    uit.to_excel(DOEL, index=False, sheet_name='Vervangen bronnen')

    wb = load_workbook(DOEL)
    ws = wb['Vervangen bronnen']
    br = {'Nr': 6, 'Oordeel': 11, 'In gebruik': 11, 'Vraag NL': 56, 'Antwoord': 12,
          'Dode link (oud)': 48, 'Nieuwe bron': 48, 'Toelichting': 70, 'Jouw besluit': 18}
    for i, k in enumerate(uit.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = br.get(k, 18)
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    i_o = list(uit.columns).index('Oordeel') + 1
    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(uit.columns) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=uit.columns[kol - 1] in ('Vraag NL', 'Toelichting'))
        o = ws.cell(row=rij, column=i_o).value
        ws.cell(row=rij, column=i_o).fill = PatternFill(
            'solid', fgColor='F8CBCB' if o == 'fout' else 'D6EAD6')
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(uit.columns))}{ws.max_row}'
    wb.save(DOEL)

    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_vervangen_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx'))
    boek = load_workbook(REVIEW)
    blad = boek['Vragen']
    kol = [c.value for c in blad[1]]
    i_bron = kol.index('Bron (geverifieerd)') + 1
    i_bew = kol.index('Bewijszin') + 1
    i_vert = kol.index('Vertrouwen bron') + 1
    i_st = kol.index('Status oude bron') + 1
    n = 0
    for rij in range(2, blad.max_row + 1):
        nr = blad.cell(row=rij, column=1).value
        if nr not in BRONNEN:
            continue
        bron, oordeel, opmerking = BRONNEN[nr]
        blad.cell(row=rij, column=i_bron).value = bron
        blad.cell(row=rij, column=i_bew).value = opmerking[:400]
        blad.cell(row=rij, column=i_vert).value = 'handmatig'
        blad.cell(row=rij, column=i_st).value = 'dode link vervangen'
        n += 1
    boek.save(REVIEW)

    print(uit['Oordeel'].value_counts().to_string())
    print(f'\nvervangen in het reviewblad: {n}')
    print(f'-> {DOEL}')


if __name__ == '__main__':
    main()

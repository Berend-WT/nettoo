# -*- coding: utf-8 -*-
"""Netto — voeg de bronvondsten samen tot één werkmap en schrijf ze terug.

Draaien:  python tools/maak_bronnenrapport.py

Combineert de twee controles:
  bronnen_gevonden.csv   Wikipedia-artikel + de zin waarin ons getal staat.
                         Breed bereik, tekstherkenning, dus een enkele valse
                         treffer zit erbij — vandaar dat de bewijszin meegaat.
  wikidata_controle.csv  Gestructureerde waarde bij een expliciete eigenschap.
                         Smal bereik, maar geen leesfouten mogelijk.

De geverifieerde bronnen gaan als nieuwe kolommen het reviewblad in. Ze
overschrijven "Bron (bestaand)" niet: dat blijft van de eigenaar.
"""

import os
import shutil
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BRONNEN = os.path.join(WORTEL, 'vragen', 'bronnen_gevonden.csv')
WIKIDATA = os.path.join(WORTEL, 'vragen', 'wikidata_controle.csv')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DOEL = os.path.join(WORTEL, 'vragen', 'bronnen_rapport.xlsx')

FONT = 'Arial'
KOP = PatternFill('solid', fgColor='1F3864')
KLEUR = {
    'bevestigd': 'D6EAD6', 'klopt': 'D6EAD6', 'bijna': 'E8F0DC',
    'getal gezien': 'FBE9A5', 'afwijkend': 'F8CBCB', 'WIJKT AF': 'F8CBCB',
    'ander onderwerp': 'E8E8E8', 'niet gevonden': 'FFFFFF', 'geen waarde': 'FFFFFF',
}


def opmaak(ws, kolommen, breedtes, statuskolom=None, wrap=()):
    for i, k in enumerate(kolommen, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breedtes.get(k, 18)
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
        c.fill = KOP
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    idx = (kolommen.index(statuskolom) + 1) if statuskolom in kolommen else None
    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(kolommen) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=kolommen[kol - 1] in wrap)
        if idx:
            s = ws.cell(row=rij, column=idx).value
            ws.cell(row=rij, column=idx).fill = PatternFill(
                'solid', fgColor=KLEUR.get(s, 'FFFFFF'))
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(kolommen))}{ws.max_row}'


def main():
    b = pd.read_csv(BRONNEN)
    w = pd.read_csv(WIKIDATA)

    # Blad 1: alle bronvondsten, bevestigde bovenaan.
    rang = {'bevestigd': 0, 'afwijkend': 1, 'getal gezien': 2, 'niet gevonden': 3}
    b['_r'] = b['Status'].map(rang).fillna(9)
    bronnen = b.sort_values(['_r', 'Nr']).drop(columns=['_r'])
    bronnen = bronnen[['Nr', 'Status', 'Categorie', 'Vraag NL', 'Antwoord',
                       'Bron', 'Artikel', 'Bewijszin', 'Toelichting',
                       'Bron (bestaand)']]

    # Blad 3: wat echt nagekeken moet worden.
    nakijken = []
    for _, r in b[b['Status'] == 'afwijkend'].iterrows():
        nakijken.append({'Nr': int(r['Nr']), 'Herkomst': 'Wikipedia',
                         'Vraag NL': r['Vraag NL'], 'Ons antwoord': r['Antwoord'],
                         'Wat er is gevonden': r['Bewijszin'], 'Bron': r['Bron'],
                         'Waarom': 'artikel noemt een ander getal bij dezelfde termen'})
    for _, r in w[w['Oordeel'].isin(['WIJKT AF', 'ander onderwerp'])].iterrows():
        nakijken.append({'Nr': int(r['Nr']), 'Herkomst': 'Wikidata',
                         'Vraag NL': r['Vraag NL'], 'Ons antwoord': r['Ons antwoord'],
                         'Wat er is gevonden': f"{r['Wikidata-waarde']} bij {r['Eigenschap']}",
                         'Bron': r['Bron'],
                         'Waarom': ('item hoort bij een ander onderwerp'
                                    if r['Oordeel'] == 'ander onderwerp'
                                    else f"wijkt {r['Afwijking']} af")})
    nakijken = pd.DataFrame(nakijken).sort_values(['Herkomst', 'Nr'])

    with pd.ExcelWriter(DOEL, engine='openpyxl') as xl:
        bronnen.to_excel(xl, index=False, sheet_name='Bronnen')
        w.to_excel(xl, index=False, sheet_name='Wikidata-controle')
        nakijken.to_excel(xl, index=False, sheet_name='Nakijken')

    wb = load_workbook(DOEL)
    opmaak(wb['Bronnen'], list(bronnen.columns),
           {'Nr': 6, 'Status': 14, 'Categorie': 22, 'Vraag NL': 56, 'Antwoord': 11,
            'Bron': 44, 'Artikel': 26, 'Bewijszin': 74, 'Toelichting': 32,
            'Bron (bestaand)': 30},
           'Status', wrap=('Vraag NL', 'Bewijszin'))
    opmaak(wb['Wikidata-controle'], list(w.columns),
           {'Nr': 6, 'Vraag NL': 54, 'Ons antwoord': 13, 'Eigenschap': 24,
            'Onderwerp': 22, 'Wikidata-item': 22, 'Wikidata-waarde': 16,
            'Peiljaar': 10, 'Oordeel': 15, 'Afwijking': 11, 'Bron': 38},
           'Oordeel', wrap=('Vraag NL',))
    opmaak(wb['Nakijken'], list(nakijken.columns),
           {'Nr': 6, 'Herkomst': 12, 'Vraag NL': 54, 'Ons antwoord': 13,
            'Wat er is gevonden': 70, 'Bron': 40, 'Waarom': 34},
           None, wrap=('Vraag NL', 'Wat er is gevonden'))
    wb.save(DOEL)

    # Terugschrijven naar het reviewblad, in eigen kolommen.
    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW), f'_backup_bronnen_{date.today():%Y-%m-%d}_'
        + os.path.basename(REVIEW)))
    boek = load_workbook(REVIEW)
    ws = boek['Vragen']
    kolommen = [c.value for c in ws[1]]
    for naam in ('Bron (geverifieerd)', 'Bewijszin', 'Controle'):
        if naam not in kolommen:
            kolommen.append(naam)
            ws.cell(row=1, column=len(kolommen), value=naam).font = Font(
                name=FONT, bold=True, color='FFFFFF', size=11)
            ws.cell(row=1, column=len(kolommen)).fill = KOP

    i_bron = kolommen.index('Bron (geverifieerd)') + 1
    i_bewijs = kolommen.index('Bewijszin') + 1
    i_ctrl = kolommen.index('Controle') + 1

    per_nr = {int(r['Nr']): r for _, r in b.iterrows()}
    wd = {int(r['Nr']): r for _, r in w.iterrows()}
    geschreven = 0
    for rij in range(2, ws.max_row + 1):
        nr = ws.cell(row=rij, column=1).value
        r = per_nr.get(nr)
        if r is not None and r['Status'] == 'bevestigd':
            ws.cell(row=rij, column=i_bron).value = r['Bron']
            ws.cell(row=rij, column=i_bewijs).value = str(r['Bewijszin'])[:400]
            geschreven += 1
        if nr in wd:
            oud = ws.cell(row=rij, column=i_ctrl).value or ''
            ws.cell(row=rij, column=i_ctrl).value = (
                f"Wikidata: {wd[nr]['Oordeel']}" + (f' — {oud}' if oud else ''))
    for kol in (i_bron, i_bewijs, i_ctrl):
        ws.column_dimensions[get_column_letter(kol)].width = 46
    boek.save(REVIEW)

    print(b['Status'].value_counts().to_string())
    print()
    print(f'bronnen weggeschreven naar het reviewblad : {geschreven}')
    print(f'rijen op het blad "Nakijken"              : {len(nakijken)}')
    print(f'-> {DOEL}')


if __name__ == '__main__':
    main()

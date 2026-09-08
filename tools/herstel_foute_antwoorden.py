# -*- coding: utf-8 -*-
"""Netto — corrigeer de eenentwintig foute antwoorden in de vragenbank.

Draaien:  python tools/herstel_foute_antwoorden.py [--schrijf]

Zonder --schrijf toont het script alleen wat het zou doen.

Dit script raakt alleen de vragenbank. De puzzelbestanden blijven ongemoeid;
die komen bij de puzzelherbouw aan de beurt. Wel schrijft het een overzicht weg
van elke puzzel die door een correctie ongeldig wordt, zodat die herbouw weet
wat er opgeruimd moet worden.

De fouten vallen in twee soorten, en dat verschil bepaalt de schade.

Bij acht vragen klopt het getal en deugt de vraag niet. Een potvis duikt tot
2000 meter, niet 2000 kilometer; Parmigiano rijpt twaalf maanden, niet twaalf
dagen. Daar verandert alleen de vraagtekst. Het antwoord blijft staan, dus elke
som waarin die vraag voorkomt blijft kloppen — deze acht kosten geen puzzel.

Bij twaalf vragen is het antwoord zelf fout, en bij een vervalt de vraag. Dan
breekt elke puzzel waarin die vraag zit, want de som gaat uit van het oude
getal. Zo'n puzzel is niet te lijmen door er een ander getal in te zetten: de
andere twee vragen hebben hun eigen vaste antwoord. Bij de herbouw is de keuze
per puzzel om er een andere vraag met precies het oude antwoord in te zetten, of
de puzzel te laten vervallen. Het overzicht noemt per geval of zo'n vervanger
bestaat.
"""

import json
import os
import shutil
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DATA = os.path.join(WORTEL, 'data')
RAPPORT = os.path.join(WORTEL, 'vragen', 'puzzels_om_te_herbouwen.xlsx')

# ---- 1. Vraagtekst fout, antwoord goed. Sommen blijven heel. -------------
TEKST = {
    156: 'Hoeveel meter diep kan de potvis duiken?',
    491: 'Hoeveel vakjes heeft een standaard Mens-erger-je-niet-bord?',
    599: 'Hoeveel zijvlakken heeft een icosaeder?',
    706: 'Hoeveel maanden rijpt Parmigiano Reggiano minimaal?',
    783: 'Hoeveel kilometer staaldraad zit er in beide hoofdkabels van de Golden Gate Bridge samen?',
    817: 'Hoeveel ton ijzer werd gebruikt voor de Eiffeltoren?',
    826: 'Hoeveel bovengrondse verdiepingen heeft het Witte Huis?',
    1189: 'Hoeveel meter lang is het hoofdspansegment van de Van Brienenoordbrug?',
}

# ---- 2. Antwoord fout. Puzzels die de vraag gebruiken breken. ------------
ANTWOORD = {
    42: (1665, 1666, 'De volkstelling van Jean Talon begon op 21 maart 1666.'),
    69: (255, 246, 'Zeng Jinlian, de langste vrouw ooit, was 246,3 centimeter.'),
    72: (97, 100, 'Lucky Diamond Rich is voor honderd procent bedekt.'),
    93: (1892, 1887, 'De eerste internationale telefoonlijn opende op 24 februari 1887 tussen '
                     'Parijs en Brussel. De onderzeese verbinding Londen-Parijs volgde in 1891.'),
    169: (21, 11, 'Elf landen voerden de euro in 1999 in.'),
    358: (152, 147, 'De draagtijd van een schaap is gemiddeld 147 dagen.'),
    404: (75, 74, 'Windkracht 8 loopt tot 74 kilometer per uur; 75 hoort al bij windkracht 9.'),
    438: (36, 45, 'Thumbelina, het kleinste paard ooit, was 44,5 centimeter; afgerond 45.'),
    703: (41198, 41822, 'Het jaarverslag geeft 41.822 restaurants eind 2023.'),
    789: (40, 38, 'Christus de Verlosser is met sokkel 38 meter hoog. Let op: bijna-duplicaat '
                  'van vraag 787.'),
    988: (4, 5, 'Vijf van de zeven wereldwonderen lagen in Azie of Noord-Afrika.'),
    1089: (12500, 13000, 'Lidl passeerde dertienduizend filialen.'),
}

# ---- 3. Vraag vervalt. -------------------------------------------------
VERVALT = {
    1457: 'Vraagt hoeveel megabyte in "precies een gigabyte" gaan; dat zijn er 1000, niet 1024. '
          'Vraag 1458 stelt dezelfde vraag mét de binaire voorwaarde en is daarmee correct. '
          'Deze vraag is een dubbele met een fout antwoord.',
}

BESTANDEN = ['netto_frontend_puzzles.js', 'netto_breinkrakers.js',
             'netto_race_sets.js', 'netto_race_pool.js']


def laad(pad):
    t = open(pad, encoding='utf-8').read()
    i = t.index('=', t.index('window.')) + 1
    return json.loads(t[i:].strip().rstrip(';'))


def vragen_in(puzzel):
    """Geeft (label, antwoord) voor elke vraag in een puzzel.

    Twee vormen komen voor: platte velden q1_label/q1_answer in de puzzel- en
    racebestanden, en een genest object q1 = {label, answer, category} in de
    breinkrakers.
    """
    uit = []
    for i in (1, 2, 3, 4):
        genest = puzzel.get(f'q{i}')
        if isinstance(genest, dict) and 'label' in genest:
            uit.append((genest['label'], genest.get('answer')))
        elif f'q{i}_label' in puzzel:
            uit.append((puzzel[f'q{i}_label'], puzzel.get(f'q{i}_answer')))
    return uit


def alle_puzzels(data):
    if isinstance(data, list):
        return [('', p) for p in data]
    uit = []
    for sleutel, lijst in data.items():
        if isinstance(lijst, list):
            uit += [(sleutel, p) for p in lijst if isinstance(p, dict)]
    return uit


def main():
    schrijf = '--schrijf' in sys.argv
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    vraag_van = {int(r['Nr']): str(r['Vraag NL']) for _, r in d.iterrows()}

    # Kandidaat-vervangers: vragen met een deugdelijk antwoord, op waarde gegroepeerd.
    onbruikbaar = set(d.loc[d['Vertrouwen bron'].isin(
        ['antwoord fout', 'onbruikbaar']), 'Nr']) | set(VERVALT)
    per_waarde = {}
    for _, r in d.iterrows():
        if r['Nr'] in onbruikbaar:
            continue
        try:
            a = int(r['Antwoord'])
        except (TypeError, ValueError):
            continue
        per_waarde.setdefault(a, []).append(str(r['Vraag NL']))

    breekt = {}
    for nr, (oud, nieuw, _) in ANTWOORD.items():
        breekt[vraag_van[nr]] = (nr, oud, len(per_waarde.get(oud, [])))
    for nr in VERVALT:
        breekt[vraag_van[nr]] = (nr, None, 0)

    rijen = []
    for naam in BESTANDEN:
        data = laad(os.path.join(DATA, naam))
        for sleutel, p in alle_puzzels(data):
            for label, _ in vragen_in(p):
                if label not in breekt:
                    continue
                nr, oud, aantal = breekt[label]
                rijen.append({
                    'Bestand': naam.replace('netto_', '').replace('.js', ''),
                    'Lijst': sleutel, 'Puzzel': p.get('id'),
                    'Datum': p.get('date', ''),
                    'Som': p.get('calculation') or p.get('formula', ''),
                    'Vraag Nr': nr, 'Vraag': label,
                    'Oud antwoord': oud,
                    'Vervangers met dat antwoord': aantal,
                    'Advies': 'vervangen' if aantal else 'laten vervallen',
                })

    r = pd.DataFrame(rijen)
    print('Vraagtekst gecorrigeerd (antwoord blijft, geen puzzel raakt stuk):')
    for nr in sorted(TEKST):
        print(f'  {nr:5d}  {TEKST[nr][:78]}')
    print(f'\nAntwoord gecorrigeerd: {len(ANTWOORD)}   vraag vervallen: {len(VERVALT)}')
    print(f'\nPuzzels die daardoor ongeldig worden: {len(r)}')
    if len(r):
        print(r.groupby(['Bestand', 'Advies']).size().to_string())

    if not schrijf:
        print('\n(proefdraai — voeg --schrijf toe om op te slaan)')
        return

    # ---- overzicht wegschrijven -----------------------------------------
    r = r.sort_values(['Bestand', 'Lijst', 'Puzzel'])
    r.to_excel(RAPPORT, index=False, sheet_name='Te herbouwen')
    wb = load_workbook(RAPPORT)
    ws = wb['Te herbouwen']
    br = {'Bestand': 18, 'Lijst': 16, 'Puzzel': 16, 'Datum': 12, 'Som': 30,
          'Vraag Nr': 9, 'Vraag': 62, 'Oud antwoord': 13,
          'Vervangers met dat antwoord': 13, 'Advies': 18}
    for i, k in enumerate(r.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = br.get(k, 16)
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    i_a = list(r.columns).index('Advies') + 1
    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(r.columns) + 1):
            ws.cell(row=rij, column=kol).font = Font(name='Arial', size=10)
            ws.cell(row=rij, column=kol).alignment = Alignment(
                vertical='top', wrap_text=r.columns[kol - 1] == 'Vraag')
        a = ws.cell(row=rij, column=i_a).value
        ws.cell(row=rij, column=i_a).fill = PatternFill(
            'solid', fgColor='D6EAD6' if a == 'vervangen' else 'F8CBCB')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(r.columns))}{ws.max_row}'
    wb.save(RAPPORT)

    # ---- vragenbank bijwerken -------------------------------------------
    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_antwoorden_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx'))
    boek = load_workbook(REVIEW)
    blad = boek['Vragen']
    kol = [c.value for c in blad[1]]
    i_vraag = kol.index('Vraag NL') + 1
    i_antw = kol.index('Antwoord') + 1
    i_vert = kol.index('Vertrouwen bron') + 1
    i_st = kol.index('Status oude bron') + 1
    i_let = kol.index('Let op') + 1

    for rij in range(2, blad.max_row + 1):
        nr = blad.cell(row=rij, column=1).value
        if nr in TEKST:
            oud = blad.cell(row=rij, column=i_vraag).value
            blad.cell(row=rij, column=i_vraag).value = TEKST[nr]
            blad.cell(row=rij, column=i_vert).value = 'handmatig'
            blad.cell(row=rij, column=i_st).value = 'vraagtekst gecorrigeerd'
            blad.cell(row=rij, column=i_let).value = f'was: {oud}'[:250]
        elif nr in ANTWOORD:
            oud, nieuw, waarom = ANTWOORD[nr]
            blad.cell(row=rij, column=i_antw).value = nieuw
            blad.cell(row=rij, column=i_vert).value = 'handmatig'
            blad.cell(row=rij, column=i_st).value = 'antwoord gecorrigeerd'
            blad.cell(row=rij, column=i_let).value = f'was {oud}. {waarom}'[:250]
        elif nr in VERVALT:
            blad.cell(row=rij, column=i_vert).value = 'onbruikbaar'
            blad.cell(row=rij, column=i_st).value = 'vraag vervallen'
            blad.cell(row=rij, column=i_let).value = VERVALT[nr][:250]
    boek.save(REVIEW)
    print(f'\nvragenbank bijgewerkt\n-> {RAPPORT}')


if __name__ == '__main__':
    main()

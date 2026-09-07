# -*- coding: utf-8 -*-
"""Netto — zet de categoriecorrecties door in de vragenbank en het reviewblad.

Draaien:  python tools/pas_categorieen_toe.py

Raakt drie bestanden aan:
  - vragen/1000+ vragen netjes gecategoriseerd.xlsx   (de bank zelf)
  - vragen/vragen_review_compleet.xlsx                (het reviewblad)
  - de kopie van het reviewblad in OneDrive

Er wordt altijd eerst een backup weggeschreven naast het origineel.
"""

import os
import shutil
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from categorie_correcties import (
    KAPOTTE_VRAGEN,
    OPGEHEVEN_CATEGORIEEN,
    SPELFOUTEN,
    alle_correcties,
)

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK = os.path.join(WORTEL, 'vragen', '1000+ vragen netjes gecategoriseerd.xlsx')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
ONEDRIVE = os.path.join(
    os.path.expanduser('~'), 'OneDrive - Driestar-Wartburg', 'vragen_review_compleet.xlsx'
)

FONT = 'Arial'
KOP_VULLING = PatternFill('solid', fgColor='1F3864')
GEWIJZIGD = PatternFill('solid', fgColor='FFF2CC')   # zacht geel
KAPOT = PatternFill('solid', fgColor='FCE4E4')       # zacht rood


def backup(pad):
    if not os.path.exists(pad):
        return None
    doel = os.path.join(
        os.path.dirname(pad),
        f'_backup_{date.today():%Y-%m-%d}_{os.path.basename(pad)}',
    )
    shutil.copy2(pad, doel)
    return doel


def pas_toe(df):
    """Zet de nieuwe categorie in df en geeft terug wat er veranderde."""
    correcties = alle_correcties()
    # Bij een tweede run mag "was" niet de al gecorrigeerde waarde worden:
    # de oorspronkelijke categorie is dan de kolom die er al staat.
    if 'Categorie (was)' not in df.columns:
        df['Categorie (was)'] = df['Categorie'].astype(str).str.strip()
    nieuw = []
    for _, r in df.iterrows():
        nr = int(r['Nr'])
        oud = str(r['Categorie']).strip()
        nieuw.append(correcties.get(nr, oud))
    df['Categorie'] = nieuw

    # Elke overgebleven vraag uit een opgeheven categorie is een gat in de
    # correctielijst; die willen we hard zien in plaats van stil doorlaten.
    rest = df[df['Categorie'].isin(OPGEHEVEN_CATEGORIEEN)]
    if len(rest):
        raise SystemExit(
            f'{len(rest)} vragen staan nog in een opgeheven categorie: '
            + ', '.join(str(int(n)) for n in rest['Nr'].head(20))
        )

    df['Advies Claude'] = [
        KAPOTTE_VRAGEN.get(int(r['Nr']), '') for _, r in df.iterrows()
    ]
    for nr, (fout, goed) in SPELFOUTEN.items():
        rij = df.index[df['Nr'] == nr]
        if len(rij):
            i = rij[0]
            df.at[i, 'Vraag NL'] = str(df.at[i, 'Vraag NL']).replace(fout, goed)
    return df


def schrijf(df, pad, titel):
    kolommen = [
        'Nr', 'In gebruik', 'Categorie', 'Categorie (was)', 'Vraag NL',
        'Vraag EN (zoekhulp)', 'Antwoord', 'Bron (bestaand)',
        'Bron (jouw aanvulling)', 'Commons-fotolink', 'Verwijderen',
        'Advies Claude', 'Jouw opmerking', 'Let op',
    ]
    kolommen = [k for k in kolommen if k in df.columns]
    df = df[kolommen]

    with pd.ExcelWriter(pad, engine='openpyxl') as w:
        df.to_excel(w, index=False, sheet_name='Vragen')

    wb = load_workbook(pad)
    ws = wb['Vragen']
    ws.title = 'Vragen'

    breedtes = {
        'Nr': 6, 'In gebruik': 10, 'Categorie': 26, 'Categorie (was)': 26,
        'Vraag NL': 78, 'Vraag EN (zoekhulp)': 62, 'Antwoord': 14,
        'Bron (bestaand)': 40, 'Bron (jouw aanvulling)': 34,
        'Commons-fotolink': 30, 'Verwijderen': 12, 'Advies Claude': 62,
        'Jouw opmerking': 28, 'Let op': 30,
    }
    for i, kol in enumerate(kolommen, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breedtes.get(kol, 20)
        cel = ws.cell(row=1, column=i)
        cel.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
        cel.fill = KOP_VULLING
        cel.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30

    i_cat = kolommen.index('Categorie') + 1
    i_was = kolommen.index('Categorie (was)') + 1
    i_adv = kolommen.index('Advies Claude') + 1

    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(kolommen) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical='top', wrap_text=(kol in (5, 6, i_adv)))
        if ws.cell(row=rij, column=i_cat).value != ws.cell(row=rij, column=i_was).value:
            ws.cell(row=rij, column=i_cat).fill = GEWIJZIGD
            ws.cell(row=rij, column=i_cat).font = Font(name=FONT, size=10, bold=True)
        if ws.cell(row=rij, column=i_adv).value:
            ws.cell(row=rij, column=i_adv).fill = KAPOT

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(kolommen))}{ws.max_row}'

    # Leeswijzer als eerste tabblad, zodat duidelijk is wat er veranderd is.
    uitleg = wb.create_sheet('Leeswijzer', 0)
    regels = [
        (titel, True),
        ('', False),
        ('Geel gemarkeerde categorie = door mij gewijzigd. De oude staat ernaast in "Categorie (was)".', False),
        ('Rood gemarkeerd "Advies Claude" = de vraag is inhoudelijk stuk. Mijn advies is schrappen;', False),
        ('de reden staat erbij. Zet zelf "ja" in de kolom "Verwijderen" als je het ermee eens bent.', False),
        ('', False),
        ('Twee categorieën zijn opgeheven:', True),
        ('  Topografie  — beschreef hetzelfde als Geografie ("Hoeveel landen grenzen aan X?" stond in allebei).', False),
        ('  Nederlands  — was een land, geen onderwerp. Die vragen hebben nu hun echte onderwerp gekregen.', False),
        ('                Nederland blijft gewoon bestaan als race-set in het spel.', False),
        ('', False),
        ('Waarom dit ertoe doet: de categorie bepaalt welk icoon op de landingspagina onder de daily', False),
        ('verschijnt. Een honkbalvraag onder "Dieren" tekende een pootafdruk boven een vraag over honkbal.', False),
        ('', False),
        ('Nog voor jou: kolom "Bron (jouw aanvulling)" en "Commons-fotolink" zijn nog leeg.', False),
    ]
    for r, (tekst, vet) in enumerate(regels, start=1):
        c = uitleg.cell(row=r, column=1, value=tekst)
        c.font = Font(name=FONT, size=11, bold=vet)
    uitleg.column_dimensions['A'].width = 118

    wb.save(pad)


def main():
    correcties = alle_correcties()
    df = pd.read_excel(REVIEW).sort_values('Nr').reset_index(drop=True)
    df = pas_toe(df)

    gewijzigd = int((df['Categorie'] != df['Categorie (was)']).sum())
    kapot = int((df['Advies Claude'] != '').sum())

    for pad in (REVIEW, BANK, ONEDRIVE):
        b = backup(pad)
        if b:
            print(f'backup -> {os.path.basename(b)}')

    schrijf(df, REVIEW, 'Netto — vragenbank, categorieën herzien')
    try:
        shutil.copy2(REVIEW, ONEDRIVE)
        print(f'onedrive   -> {ONEDRIVE}')
    except PermissionError:
        # Staat open in Excel. Het reviewblad in de repo is bijgewerkt; de
        # kopie is een gemak, geen voorwaarde.
        print('LET OP: OneDrive-kopie staat open in Excel en is NIET bijgewerkt.')

    # De bank houdt zijn eigen, smallere kolomindeling. Hij heeft geen
    # Nr-kolom, dus koppelen gaat op de vraagtekst zoals die vóór de
    # spellingcorrectie luidde.
    origineel = pd.read_excel(os.path.join(
        os.path.dirname(BANK), f'_backup_{date.today():%Y-%m-%d}_{os.path.basename(BANK)}'
    ))
    op_oude_tekst = {}
    for nr, oude_vraag in zip(df['Nr'], pd.read_excel(
        os.path.join(os.path.dirname(REVIEW),
                     f'_backup_{date.today():%Y-%m-%d}_{os.path.basename(REVIEW)}')
    ).sort_values('Nr')['Vraag NL']):
        op_oude_tekst[str(oude_vraag).strip()] = nr

    nieuwe_cat = dict(zip(df['Nr'], df['Categorie']))
    nieuwe_tekst = dict(zip(df['Nr'], df['Vraag NL']))

    kolom_cat, kolom_vraag, geraakt = [], [], 0
    for _, r in origineel.iterrows():
        nr = op_oude_tekst.get(str(r['Vraag NL']).strip())
        if nr is None:
            kolom_cat.append(str(r['Categorie']).strip())
            kolom_vraag.append(r['Vraag NL'])
        else:
            kolom_cat.append(nieuwe_cat[nr])
            kolom_vraag.append(nieuwe_tekst[nr])
            geraakt += 1

    bank = origineel.copy()
    bank['Categorie'] = kolom_cat
    bank['Vraag NL'] = kolom_vraag
    bank.to_excel(BANK, index=False)

    print(f'\ncategorieën gewijzigd : {gewijzigd} van {len(df)}')
    print(f'vragen als stuk gemerkt: {kapot}')
    print(f'bank gekoppeld op tekst: {geraakt} van {len(origineel)}')
    print(f'\nreviewblad -> {REVIEW}')


if __name__ == '__main__':
    main()

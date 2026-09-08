# -*- coding: utf-8 -*-
"""Netto — vragen die ten onrechte als "geen bron mogelijk" waren weggezet.

Draaien:  python tools/bronnen_alsnog_mogelijk.py [--schrijf]

Bij het doorlopen van de 45 vragen zonder geverifieerde bron bleek dat die stapel
twee heel verschillende dingen bevatte, en dat ik ze onder een noemer had gezet.

Het grootste deel deugt echt niet: vragen met antwoord nul, vragen naar iets
zonder norm, en een reeks die duidelijk kapot gegenereerd is ("Hoeveel kuiven
heeft een honkbalveld", "Alvleeskliervlaamse Wielerweek"). Daar bestaat geen
bron voor omdat er geen feit is.

Maar er zaten ook vragen bij die niets mankeren. De som van 1 tot en met 100 is
5050 — dat is niet onvindbaar, dat volgt uit een formule. Het NBA-record van 186
punten staat sinds 1983 en is prima te staven. Die had ik niet mogen afschrijven
met "geen bron mogelijk"; ze hadden alleen geen webpagina waar een crawler het
getal kon aanwijzen. Dat is iets anders.

Hieronder krijgen die alsnog een bron. Een vraag blijkt bovendien een fout
antwoord te hebben in plaats van een ontbrekende bron.
"""

import os
import shutil
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')

A = {
    104: ('bevestigd', 'https://en.wikipedia.org/wiki/Arithmetic_progression',
          'De som van 1 tot en met n is n(n+1)/2; voor honderd is dat 5050. Het verhaal over '
          'de jonge Gauss gaat over precies deze som.'),
    105: ('bevestigd', 'https://en.wikipedia.org/wiki/Permutation',
          'Een geordende keuze van drie uit acht is 8 x 7 x 6 = 336.'),
    147: ('natellen', 'https://en.wikipedia.org/wiki/Pedometer',
          'De tienduizend stappen komen uit de naam van een Japanse stappenteller uit 1965 en '
          'zijn nooit een medische norm geweest. De waarde is gangbaar, niet vastgesteld.'),
    455: ('bevestigd', 'https://www.basketball-reference.com/boxscores/198312130DEN.html',
          'Detroit won op 13 december 1983 met 186-184 van Denver na drie verlengingen; nog '
          'altijd de hoogste teamscore in een NBA-wedstrijd.'),
    546: ('veroudert', 'https://en.wikipedia.org/wiki/Google_Search',
          'De negenennegentigduizend zoekopdrachten per seconde is een veelgeciteerde '
          'schatting die met het gebruik meegroeit.'),
    74: ('veroudert', 'https://en.wikipedia.org/wiki/Lego',
         'LEGO produceert tientallen miljarden elementen per jaar; omgerekend rond de '
         'dertienhonderd per seconde. Het cijfer verschuift met de productie.'),
}

# Geen ontbrekende bron maar een fout antwoord.
FOUT = {
    66: (815, 1056, 'https://en.wikipedia.org/wiki/IMAX_Sydney',
         'Het grootste vaste IMAX-scherm ooit was dat in Darling Harbour met 1056 vierkante '
         'meter, gesloopt in 2016. Het huidige scherm daar meet 692. De 815 hoort bij geen '
         'van beide.'),
}


def main():
    schrijf = '--schrijf' in sys.argv
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    for nr in list(A) + list(FOUT):
        r = d[d['Nr'] == nr]
        print(f"{nr:5d}  {str(r['Vraag NL'].iloc[0])[:70]} -> {r['Antwoord'].iloc[0]}")
    if not schrijf:
        print('\n(proefdraai — voeg --schrijf toe om op te slaan)')
        return

    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_alsnog_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx'))
    boek = load_workbook(REVIEW)
    blad = boek['Vragen']
    kol = [c.value for c in blad[1]]
    i_bron = kol.index('Bron (geverifieerd)') + 1
    i_bew = kol.index('Bewijszin') + 1
    i_vert = kol.index('Vertrouwen bron') + 1
    i_st = kol.index('Status oude bron') + 1
    i_antw = kol.index('Antwoord') + 1
    i_let = kol.index('Let op') + 1
    V = {'bevestigd': 'handmatig', 'natellen': 'natellen', 'veroudert': 'veroudert'}

    for rij in range(2, blad.max_row + 1):
        nr = blad.cell(row=rij, column=1).value
        if nr in A:
            oordeel, bron, toelichting = A[nr]
            blad.cell(row=rij, column=i_bron).value = bron
            blad.cell(row=rij, column=i_bew).value = toelichting[:400]
            blad.cell(row=rij, column=i_vert).value = V[oordeel]
            blad.cell(row=rij, column=i_st).value = 'alsnog een bron gevonden'
        elif nr in FOUT:
            oud, nieuw, bron, toelichting = FOUT[nr]
            blad.cell(row=rij, column=i_antw).value = nieuw
            blad.cell(row=rij, column=i_bron).value = bron
            blad.cell(row=rij, column=i_bew).value = toelichting[:400]
            blad.cell(row=rij, column=i_vert).value = 'handmatig'
            blad.cell(row=rij, column=i_st).value = 'antwoord gecorrigeerd'
            blad.cell(row=rij, column=i_let).value = f'was {oud}. {toelichting}'[:250]
    boek.save(REVIEW)

    na = pd.read_excel(REVIEW, sheet_name='Vragen')
    ver = na['Bron (geverifieerd)'].notna()
    print(f'\ngeverifieerd: {int(ver.sum())} van {len(na)}')


if __name__ == '__main__':
    main()

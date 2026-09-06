#!/usr/bin/env python3
"""Verwerkt het ingevulde reviewblad van de vragenbank.

Doet drie dingen met wat de redacteur heeft ingevuld:

1. Bronnen terugschrijven naar de vragenbank.
2. Fotolinks ophalen van Wikimedia Commons: licentie controleren, afbeelding
   downloaden en de bronvermelding vastleggen.
3. Verwijder-markeringen verwerken, met een waarschuwing als een vraag nog in
   een puzzel zit.

De vragenbank wordt niet blind overschreven: er komt eerst een backup naast.
Foto's worden niet in de database gezet — dat kan dit script niet, want daar is
een admin-sessie voor nodig. In plaats daarvan komt er SQL uit die de eigenaar
zelf kan draaien.

Gebruik:
    python tools/verwerk_vragen_review.py --review pad/naar/ingevuld.xlsx
    python tools/verwerk_vragen_review.py --review ... --dry-run
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import urllib.parse
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "fotos"))
import maak_fotocatalogus as catalogus  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
BANK = ROOT / "vragen" / "1000+ vragen netjes gecategoriseerd.xlsx"
PUZZELS = ROOT / "netto_frontend_puzzles.js"
ASSET_DIR = ROOT / "fotos" / "assets" / "vragen"
SQL_OUT = ROOT / "supabase" / "foto_toewijzingen.sql"

# Kolomindexen in het reviewblad (0-based), zie tools/maak_vragen_reviewblad.py
K_VRAAG_NL, K_ANTWOORD = 3, 5
K_BRON_NIEUW, K_FOTOLINK, K_VERWIJDER, K_OPMERKING = 7, 8, 9, 10


def commons_titel(url: str) -> str | None:
    """Haalt de bestandstitel uit een Commons-bestandspagina-URL."""
    if not url or "commons.wikimedia.org" not in url:
        return None
    pad = urllib.parse.urlparse(url).path
    match = re.search(r"/wiki/(File:.+|Bestand:.+)$", pad)
    if not match:
        return None
    titel = urllib.parse.unquote(match.group(1)).replace("_", " ")
    return "File:" + titel.split(":", 1)[1]


def haal_commons_bestand(titel: str) -> dict:
    """Metadata en licentie van één Commons-bestand.

    Weigert alles wat niet herbruikbaar is; dat is precies de controle die de
    redacteur niet zelf hoeft te doen.
    """
    data = catalogus.request_json(catalogus.COMMONS_API, {
        "action": "query", "format": "json", "formatversion": "2",
        "titles": titel, "prop": "imageinfo",
        "iiprop": "url|extmetadata", "iiurlwidth": "1200",
    })
    pages = (data.get("query") or {}).get("pages") or []
    if not pages or pages[0].get("missing"):
        raise RuntimeError("bestand bestaat niet op Commons")
    info = (pages[0].get("imageinfo") or [{}])[0]
    meta = info.get("extmetadata") or {}

    licentie = catalogus.commons_license(meta)
    if not licentie:
        naam = catalogus.metadata_value(meta, "LicenseShortName") or "onbekend"
        raise RuntimeError(f"licentie niet herbruikbaar: {naam}")
    licentienaam, _ = licentie

    maker = catalogus.clean_html(catalogus.metadata_value(meta, "Artist")) or "onbekende maker"
    omschrijving = catalogus.clean_html(catalogus.metadata_value(meta, "ImageDescription"))
    schone_titel = titel.removeprefix("File:").removesuffix(".jpg").removesuffix(".png")
    return {
        "titel": schone_titel,
        "maker": maker,
        "licentie": licentienaam,
        "alt": (omschrijving or schone_titel)[:160],
        "credit": f"{schone_titel} · {maker} · {licentienaam}",
        "bron": f"https://commons.wikimedia.org/wiki/{urllib.parse.quote(titel.replace(' ', '_'))}",
        "download": info.get("thumburl") or info.get("url"),
    }


def laad_vraag_naar_daily() -> dict[str, list[str]]:
    """Per vraagtekst de datums van de dailies waarin hij voorkomt."""
    src = PUZZELS.read_text(encoding="utf-8")
    data = json.loads(src[src.index("=") + 1:].rstrip().rstrip(";"))
    kaart: dict[str, list[str]] = {}
    for puzzel in data.get("daily") or []:
        for slot in ("q1", "q2", "q3"):
            vraag = (puzzel.get(f"{slot}_label") or "").strip()
            if vraag and puzzel.get("date"):
                kaart.setdefault(vraag, []).append(puzzel["date"])
    return kaart


def sql_str(value) -> str:
    return "null" if value in (None, "") else "'" + str(value).replace("'", "''") + "'"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--review", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true",
                        help="Alleen rapporteren, niets schrijven of downloaden")
    args = parser.parse_args()

    ws = openpyxl.load_workbook(args.review, read_only=True)["Vragen"]
    rijen = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[K_VRAAG_NL]]

    bronnen: dict[str, str] = {}
    fotos: dict[str, dict] = {}
    verwijderen: list[str] = []
    problemen: list[str] = []

    for r in rijen:
        vraag = str(r[K_VRAAG_NL]).strip()
        if r[K_BRON_NIEUW]:
            bronnen[vraag] = str(r[K_BRON_NIEUW]).strip()
        if str(r[K_VERWIJDER] or "").strip().lower() in ("ja", "j", "x", "yes"):
            verwijderen.append(vraag)
        link = str(r[K_FOTOLINK] or "").strip()
        if not link:
            continue
        titel = commons_titel(link)
        if not titel:
            problemen.append(f"geen Commons-bestandspagina: {link[:60]} (vraag: {vraag[:40]})")
            continue
        if args.dry_run:
            fotos[vraag] = {"titel": titel}
            continue
        try:
            gegevens = haal_commons_bestand(titel)
            ASSET_DIR.mkdir(parents=True, exist_ok=True)
            veilig = re.sub(r"[^a-z0-9]+", "-", gegevens["titel"].casefold()).strip("-")[:60]
            doel = ASSET_DIR / f"{veilig or 'commons'}.jpg"
            catalogus.download(gegevens["download"], doel)
            catalogus.optimize_image(doel)
            gegevens["bestand"] = str(doel.relative_to(ROOT)).replace("\\", "/")
            fotos[vraag] = gegevens
        except Exception as err:
            problemen.append(f"{titel}: {err} (vraag: {vraag[:40]})")

    # --- verwijderingen tegen het puzzelgebruik houden ---
    vraag_naar_daily = laad_vraag_naar_daily()
    src = PUZZELS.read_text(encoding="utf-8")
    alle_puzzeldata = json.loads(src[src.index("=") + 1:].rstrip().rstrip(";"))
    in_gebruik = set()
    for naam in ("daily", "library", "reserve"):
        for puzzel in alle_puzzeldata.get(naam) or []:
            for slot in ("q1", "q2", "q3"):
                if puzzel.get(f"{slot}_label"):
                    in_gebruik.add(puzzel[f"{slot}_label"].strip())
    geblokkeerd = [v for v in verwijderen if v in in_gebruik]
    verwijderbaar = [v for v in verwijderen if v not in in_gebruik]

    print(f"Gelezen: {len(rijen)} rijen")
    print(f"  bronnen ingevuld : {len(bronnen)}")
    print(f"  fotolinks        : {len(fotos)} verwerkt, {len(problemen)} probleem")
    print(f"  verwijderen      : {len(verwijderbaar)} kan weg, {len(geblokkeerd)} zit nog in een puzzel")
    for p in problemen[:15]:
        print("   ! " + p)
    for v in geblokkeerd[:10]:
        print(f"   ! niet verwijderd, nog in gebruik: {v[:70]}")

    if args.dry_run:
        print("\n(dry-run: niets geschreven)")
        return

    # --- vragenbank bijwerken ---
    if bronnen or verwijderbaar:
        backup = BANK.with_name(f"_backup_{date.today()}_{BANK.name}")
        if not backup.exists():
            shutil.copy2(BANK, backup)
        wb = openpyxl.load_workbook(BANK)
        blad = wb["Alle vragen"]
        kop = [c.value for c in blad[1]]
        kol_vraag = kop.index("Vraag NL") + 1
        kol_bron = kop.index("Bron EN") + 1
        te_wissen = []
        bijgewerkt = 0
        for rij in range(2, blad.max_row + 1):
            vraag = str(blad.cell(rij, kol_vraag).value or "").strip()
            if vraag in bronnen:
                blad.cell(rij, kol_bron).value = bronnen[vraag]
                bijgewerkt += 1
            if vraag in verwijderbaar:
                te_wissen.append(rij)
        for rij in sorted(te_wissen, reverse=True):
            blad.delete_rows(rij, 1)
        wb.save(BANK)
        print(f"\nVragenbank bijgewerkt: {bijgewerkt} bronnen, {len(te_wissen)} vragen verwijderd")
        print(f"Backup: {backup.name}")

    # --- SQL voor de foto's ---
    if fotos:
        regels = [
            "-- Netto — foto's toewijzen aan dailies",
            "-- GEGENEREERD door tools/verwerk_vragen_review.py; niet met de hand bijwerken.",
            "--",
            "-- Alleen foto's bij vragen die in een daily zitten kunnen worden toegewezen:",
            "-- image_path hangt aan de puzzelrij, niet aan de vraag. Een foto bij een vraag",
            "-- die nergens gebruikt wordt, staat klaar maar heeft nog geen plek.",
            "",
        ]
        toegewezen = zonder_daily = 0
        for vraag, gegevens in fotos.items():
            datums = vraag_naar_daily.get(vraag) or []
            if not datums:
                zonder_daily += 1
                continue
            for datum in datums:
                regels.append(
                    "update public.puzzles set image_path = {p}, image_alt = {a}, "
                    "image_credit = {c}, image_source_url = {s} where scheduled_date = {d};".format(
                        p=sql_str(gegevens["bestand"]), a=sql_str(gegevens["alt"]),
                        c=sql_str(gegevens["credit"]), s=sql_str(gegevens["bron"]),
                        d=sql_str(datum),
                    )
                )
                toegewezen += 1
        SQL_OUT.write_text("\n".join(regels) + "\n", encoding="utf-8")
        print(f"SQL geschreven: {SQL_OUT.relative_to(ROOT)} "
              f"({toegewezen} toewijzingen, {zonder_daily} foto's bij ongebruikte vragen)")
        print("Draai daarna: python tools/sync_website.py")


if __name__ == "__main__":
    main()

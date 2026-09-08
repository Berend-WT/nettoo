# -*- coding: utf-8 -*-
"""Netto — vul de kolom "Bron (geverifieerd)" voor bronnen die al bevestigd waren.

Draaien:  python tools/backfill_bevestigde_oude_bronnen.py

Bij de controle van de 496 bestaande bronnen kregen 65 vragen de status
"bevestigd" of "getal gezien op de pagina": de opgegeven pagina is opgehaald en
het antwoord stond erop. Die uitkomst is toen alleen in "Status oude bron"
geland en niet in "Bron (geverifieerd)", waardoor ze in elke telling meetelden
als onbevestigd. Dit script schrijft de bestaande bron alsnog door.

Er wordt niets opnieuw beoordeeld — enkel een uitkomst overgezet die er al was.
"""

import os
import shutil
from datetime import date

from openpyxl import load_workbook

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')

BEVESTIGD = {'bevestigd', 'getal gezien op de pagina'}


def main():
    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_backfill_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx'))

    boek = load_workbook(REVIEW)
    blad = boek['Vragen']
    kol = [c.value for c in blad[1]]
    i_oud = kol.index('Bron (bestaand)') + 1
    i_ver = kol.index('Bron (geverifieerd)') + 1
    i_bew = kol.index('Bewijszin') + 1
    i_vert = kol.index('Vertrouwen bron') + 1
    i_st = kol.index('Status oude bron') + 1

    n = 0
    for rij in range(2, blad.max_row + 1):
        if blad.cell(row=rij, column=i_ver).value:
            continue
        if (blad.cell(row=rij, column=i_st).value or '').strip() not in BEVESTIGD:
            continue
        bron = blad.cell(row=rij, column=i_oud).value
        if not bron:
            continue
        blad.cell(row=rij, column=i_ver).value = bron
        blad.cell(row=rij, column=i_vert).value = 'hoog'
        if not blad.cell(row=rij, column=i_bew).value:
            blad.cell(row=rij, column=i_bew).value = (
                'Bestaande bron opgehaald; het antwoord staat op de pagina.')
        n += 1

    boek.save(REVIEW)
    print(f'doorgezet: {n}')


if __name__ == '__main__':
    main()

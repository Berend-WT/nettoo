# -*- coding: utf-8 -*-
"""Netto — knip uit het grote keuzeblad de rijen die er echt toe doen.

Draaien:  python tools/maak_fotokeuze_urgent.py

WAAROM
vragen/fotokeuze.xlsx heeft ruim achthonderd rijen met bijna duizend zwevende
afbeeldingen in een enkel blad. Excel tekent die allemaal tegelijk en loopt
daarop vast.

Dat is ook niet nodig, want maar een deel vraagt om een oordeel. Een vraag
waarvan de bron een Wikipedia-artikel is heeft de infoboxfoto, en die klopt per
definitie. Overblijven de vragen waar wél kandidaten voor gevonden zijn maar
waar het spel niets toont, omdat de automatische keuze zich er niet aan durfde
te wagen. Precies daar maakt een menselijk oordeel het verschil.

Dit script kopieert die rijen naar een klein bestand en neemt de miniaturen mee
uit het bestaande werkblad, dus zonder opnieuw te downloaden.

TERUGWEG
De ingevulde keuzes gaan met tools/verwerk_fotokeuze_urgent.py terug het grote
blad in, zodat maak_fotobestand.py ze gewoon oppikt.
"""

import io
import os
import sys

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GROOT = os.path.join(WORTEL, 'vragen', 'fotokeuze.xlsx')
KLEIN = os.path.join(WORTEL, 'vragen', 'fotokeuze_urgent.xlsx')
LIVE = os.path.join(WORTEL, 'data', 'netto_fotos.js')

KOLOMMEN = ['Nr', 'In gebruik', 'Vraag', 'Antwoord',
            'Foto 1', 'Foto 2', 'Foto 3', 'Keuze', 'Commons', 'Artikel', 'Opmerking']
BREEDTES = [6, 11, 60, 10, 23, 23, 23, 9, 13, 13, 30]


def main():
    bron = openpyxl.load_workbook(GROOT)
    bl = bron['Fotokeuze']
    kop = [c.value for c in bl[1]]
    k = {naam: n for n, naam in enumerate(kop, start=1)}

    # Afbeeldingen liggen los van de cellen; hun anker vertelt bij welke rij en
    # kolom ze horen. openpyxl telt daar vanaf nul, de cellen vanaf een.
    beelden = {}
    for beeld in bl._images:
        beelden.setdefault((beeld.anchor._from.row + 1, beeld.anchor._from.col + 1), []).append(beeld)

    # De rijen die er echt toe doen: er zijn kandidaten, maar het spel toont
    # niets. Dat zijn precies de vragen waar de automatische keuze zich niet
    # aan durfde te wagen, en waar jouw oordeel het verschil maakt.
    import json
    tekst = open(LIVE, encoding='utf-8').read()
    in_het_spel = set(json.loads(tekst[tekst.index('{'):tekst.rindex('}') + 1]))
    rijen = []
    for r in range(2, bl.max_row + 1):
        vraag = bl.cell(r, k['Vraag']).value
        heeft_kandidaat = any(beelden.get((r, k[f'Foto {n}'])) for n in (1, 2, 3))
        if heeft_kandidaat and str(vraag) not in in_het_spel:
            rijen.append(r)
    print(f'{len(rijen)} vragen met kandidaten maar zonder foto in het spel')

    doel = openpyxl.Workbook()
    ws = doel.active
    ws.title = 'Fotokeuze'
    donker = PatternFill('solid', fgColor='1F3864')
    for n, (naam, breedte) in enumerate(zip(KOLOMMEN, BREEDTES), start=1):
        cel = ws.cell(1, n, naam)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = donker
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(n)].width = breedte
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = 'D2'

    geel = PatternFill('solid', fgColor='FFF2CC')
    meegenomen = 0
    for nieuw, oud in enumerate(rijen, start=2):
        for naam in KOLOMMEN:
            if naam.startswith('Foto') or naam not in k:
                continue
            origineel = bl.cell(oud, k[naam])
            cel = ws.cell(nieuw, KOLOMMEN.index(naam) + 1, origineel.value)
            cel.alignment = Alignment(vertical='center', wrap_text=(naam == 'Vraag'))
            # De zoeklinks zijn hyperlinks; zonder dit blijft alleen het woord
            # "commons" staan en kun je nergens heen klikken.
            if origineel.hyperlink is not None:
                cel.hyperlink = origineel.hyperlink.target
                cel.font = Font(color='0563C1', underline='single')
            if naam == 'Keuze':
                cel.fill = geel
        for nummer in (1, 2, 3):
            for beeld in beelden.get((oud, k[f'Foto {nummer}']), []):
                data = beeld.ref if isinstance(beeld.ref, bytes) else beeld.ref.getvalue()
                kopie = XlImage(io.BytesIO(data))
                kopie.width, kopie.height = beeld.width, beeld.height
                ws.add_image(kopie, f'{get_column_letter(4 + nummer)}{nieuw}')
                meegenomen += 1
        ws.row_dimensions[nieuw].height = bl.row_dimensions[oud].height or 107.8

    ws.auto_filter.ref = f'A1:{get_column_letter(len(KOLOMMEN))}{ws.max_row}'
    doel.save(KLEIN)
    grootte = os.path.getsize(KLEIN) / 1024
    print(f'{os.path.relpath(KLEIN, WORTEL)}: {ws.max_row - 1} rijen, '
          f'{meegenomen} miniaturen, {grootte:.0f} KB')


if __name__ == '__main__':
    main()

# -*- coding: utf-8 -*-
"""Netto — werk de kolom "In gebruik" bij vanuit de puzzelbestanden.

Draaien:  python tools/werk_ingebruik_bij.py

De kolom zegt of een vraag in een dagpuzzel, een gewone puzzel, de race of een
breinkraker zit. Hij wordt niet automatisch bijgehouden, dus na elke
puzzelronde loopt hij achter — na de laatste stonden 342 vragen wél in een
puzzel zonder zo gemarkeerd te zijn, en 119 andersom.

Dat is niet onschuldig: het fotokeuzeblad kiest zijn rijen op deze kolom. Loopt
hij achter, dan zit je foto's uit te kiezen bij vragen die niemand te zien
krijgt, terwijl de vragen die wél in het spel zitten ontbreken.

Een vraag in meerdere modi krijgt het label van de zichtbaarste: een dagpuzzel
gaat voor een gewone puzzel, die weer voor race en breinkrakers.
"""

import json
import os
import shutil
import sys
from datetime import date

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DATA = os.path.join(WORTEL, 'data')
RANG = {'daily': 0, 'puzzel': 1, 'race': 2, 'breinkraker': 3, None: 9}


def lees(naam):
    with open(os.path.join(DATA, naam), encoding='utf-8') as f:
        tekst = f.read()
    begin = min((tekst.index(c) for c in '{[' if c in tekst), default=0)
    einde = max(tekst.rindex('}') if '}' in tekst else -1,
                tekst.rindex(']') if ']' in tekst else -1)
    return json.loads(tekst[begin:einde + 1])


def main():
    gebruik = {}

    def zet(label, soort):
        if label and RANG[soort] < RANG.get(gebruik.get(label)):
            gebruik[label] = soort

    for groep, lijst in lees('netto_frontend_puzzles.js').items():
        if not isinstance(lijst, list):
            continue
        soort = 'daily' if groep == 'daily' else 'puzzel'
        for p in lijst:
            for n in (1, 2, 3):
                zet(p.get(f'q{n}_label'), soort)
    for p in lees('netto_race_pool.js'):
        for n in (1, 2, 3):
            zet(p.get(f'q{n}_label'), 'race')
    for p in lees('netto_breinkrakers.js'):
        for n in (1, 2, 3, 4):
            zet(p[f'q{n}']['label'], 'breinkraker')

    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_gebruik_{date.today():%Y-%m-%d}_{os.path.basename(REVIEW)}'))
    boek = openpyxl.load_workbook(REVIEW)
    blad = boek['Vragen']
    k = {naam: n for n, naam in enumerate([c.value for c in blad[1]], start=1)}

    veranderd = 0
    from collections import Counter
    telling = Counter()
    for rij in range(2, blad.max_row + 1):
        vraag = blad.cell(rij, k['Vraag NL']).value
        nieuw = gebruik.get(str(vraag))
        cel = blad.cell(rij, k['In gebruik'])
        if (cel.value or None) != nieuw:
            cel.value = nieuw
            veranderd += 1
        telling[nieuw or 'niet in gebruik'] += 1
    boek.save(REVIEW)

    print(f'{veranderd} rijen bijgewerkt')
    for soort, aantal in telling.most_common():
        print(f'  {soort:16} {aantal}')


if __name__ == '__main__':
    main()

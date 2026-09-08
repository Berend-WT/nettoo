# -*- coding: utf-8 -*-
"""Netto — schrap de veertig vragen die geen bruikbare bron kunnen krijgen.

Draaien:  python tools/schrap_bronloze_vragen.py [--schrijf]

Zonder --schrijf toont het script alleen wat het zou doen.

WELKE VEERTIG
Achtendertig vragen hielden na alle bronrondes niets over, en twee hadden wel
een bron maar deugen niet als vraag. Vier soorten:

  kapot gegenereerd  "Hoeveel kuiven heeft een honkbalveld", en een vraag die
                     letterlijk de correctie van de schrijver bevat.
  antwoord nul       Feitelijk juist — een djembé heeft geen snaren — maar nul
                     maakt elke deelsom stuk.
  geen norm          Daslengte, veterooggaten, breedte van een horlogeband.
                     De vraag suggereert een standaard die niet bestaat.
  onvindbaar record  Langste mouwen ooit, langste snoepstaaf.

WAT ER GEBEURT
De rijen verdwijnen uit de bank en het reviewblad, en gaan naar
vragen/geschrapte_vragen.xlsx zodat het besluit terug te vinden blijft.

WAT ER NIET GEBEURT
De puzzelbestanden blijven ongemoeid — eerst vragen, dan puzzels. Puzzels die
een geschrapte vraag gebruiken worden toegevoegd aan het overzicht
vragen/puzzels_om_te_herbouwen.xlsx, zodat de herbouw ze kent.
"""

import json
import os
import shutil
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK = os.path.join(WORTEL, 'vragen', '1000+ vragen netjes gecategoriseerd.xlsx')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
GESCHRAPT = os.path.join(WORTEL, 'vragen', 'geschrapte_vragen.xlsx')
HERBOUW = os.path.join(WORTEL, 'vragen', 'puzzels_om_te_herbouwen.xlsx')
DATA = os.path.join(WORTEL, 'data')

BESTANDEN = ['netto_frontend_puzzles.js', 'netto_breinkrakers.js',
             'netto_race_sets.js', 'netto_race_pool.js']


def laad(pad):
    t = open(pad, encoding='utf-8').read()
    i = t.index('=', t.index('window.')) + 1
    return json.loads(t[i:].strip().rstrip(';'))


def vragen_in(p):
    uit = []
    for i in (1, 2, 3, 4):
        genest = p.get(f'q{i}')
        if isinstance(genest, dict) and 'label' in genest:
            uit.append(genest['label'])
        elif f'q{i}_label' in p:
            uit.append(p[f'q{i}_label'])
    return uit


def alle_puzzels(data):
    if isinstance(data, list):
        return [('', p) for p in data]
    uit = []
    for sleutel, lijst in data.items():
        if isinstance(lijst, list):
            uit += [(sleutel, p) for p in lijst if isinstance(p, dict)]
    return uit


def opmaak(pad, blad_naam, kolommen, breedtes, wrap):
    wb = load_workbook(pad)
    ws = wb[blad_naam]
    for i, k in enumerate(kolommen, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breedtes.get(k, 18)
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(kolommen) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(vertical='top', wrap_text=kolommen[kol - 1] in wrap)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(kolommen))}{ws.max_row}'
    wb.save(pad)


def main():
    schrijf = '--schrijf' in sys.argv
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    weg = d[d['Bron (geverifieerd)'].isna() | (d['Vertrouwen bron'] == 'onbruikbaar')]
    nrs = set(int(x) for x in weg['Nr'])
    labels = {str(r['Vraag NL']): int(r['Nr']) for _, r in weg.iterrows()}
    print(f'te schrappen: {len(nrs)} vragen')
    print(weg['In gebruik'].value_counts().to_string())

    # Welke puzzels raken hierdoor een vraag kwijt?
    rijen = []
    for naam in BESTANDEN:
        data = laad(os.path.join(DATA, naam))
        for sleutel, p in alle_puzzels(data):
            for label in vragen_in(p):
                if label in labels:
                    rijen.append({
                        'Bestand': naam.replace('netto_', '').replace('.js', ''),
                        'Lijst': sleutel, 'Puzzel': p.get('id'),
                        'Datum': p.get('date', ''),
                        'Som': p.get('calculation') or p.get('formula', ''),
                        'Vraag Nr': labels[label], 'Vraag': label,
                        'Oud antwoord': None,
                        'Vervangers met dat antwoord': 0,
                        'Advies': 'vraag geschrapt, puzzel laten vervallen',
                    })
    nieuw = pd.DataFrame(rijen)
    print(f'\npuzzels die een geschrapte vraag gebruiken: {len(nieuw)}')
    if len(nieuw):
        print(nieuw.groupby('Bestand').size().to_string())

    if not schrijf:
        print('\n(proefdraai — voeg --schrijf toe om op te slaan)')
        return

    # ---- geschrapte vragen bewaren --------------------------------------
    weg.to_excel(GESCHRAPT, index=False, sheet_name='Geschrapt')
    opmaak(GESCHRAPT, 'Geschrapt', list(weg.columns),
           {'Nr': 6, 'Vraag NL': 60, 'Antwoord': 12, 'Bewijszin': 74, 'Categorie': 22},
           {'Vraag NL', 'Bewijszin', 'Let op'})

    # ---- overzicht voor de puzzelherbouw aanvullen ----------------------
    if os.path.exists(HERBOUW):
        oud = pd.read_excel(HERBOUW, sheet_name='Te herbouwen')
        samen = pd.concat([oud, nieuw], ignore_index=True)
    else:
        samen = nieuw
    samen = samen.drop_duplicates(subset=['Bestand', 'Lijst', 'Puzzel', 'Vraag Nr'])
    samen = samen.sort_values(['Bestand', 'Lijst', 'Puzzel'])
    samen.to_excel(HERBOUW, index=False, sheet_name='Te herbouwen')
    opmaak(HERBOUW, 'Te herbouwen', list(samen.columns),
           {'Bestand': 18, 'Lijst': 16, 'Puzzel': 16, 'Datum': 12, 'Som': 30,
            'Vraag Nr': 9, 'Vraag': 62, 'Advies': 34}, {'Vraag'})

    # ---- rijen uit reviewblad en bank halen -----------------------------
    for pad, kolom in ((REVIEW, 'Nr'), (BANK, None)):
        if not os.path.exists(pad):
            print(f'overgeslagen (niet gevonden): {pad}')
            continue
        shutil.copy2(pad, os.path.join(
            os.path.dirname(pad),
            f'_backup_schrap_{date.today():%Y-%m-%d}_{os.path.basename(pad)}'))
        boek = load_workbook(pad)
        for blad in boek.worksheets:
            kop = [c.value for c in blad[1]]
            if 'Vraag NL' in kop:
                i_v = kop.index('Vraag NL') + 1
            elif 'Vraag' in kop:
                i_v = kop.index('Vraag') + 1
            else:
                continue
            teweg = [rij for rij in range(2, blad.max_row + 1)
                     if blad.cell(row=rij, column=i_v).value in labels]
            for rij in sorted(teweg, reverse=True):
                blad.delete_rows(rij)
            if teweg:
                print(f'{os.path.basename(pad)} / {blad.title}: {len(teweg)} rijen verwijderd')
        boek.save(pad)

    na = pd.read_excel(REVIEW, sheet_name='Vragen')
    ver = na['Bron (geverifieerd)'].notna()
    print(f'\nvragen over: {len(na)}   geverifieerd: {int(ver.sum())}')
    print(f'-> {GESCHRAPT}')
    print(f'-> {HERBOUW}')


if __name__ == '__main__':
    main()

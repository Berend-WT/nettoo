# -*- coding: utf-8 -*-
"""Netto — geef een vraag een andere tekst, overal tegelijk.

Draaien:  python tools/hernoem_vraag.py

WAAROM DIT EEN SCRIPT IS EN GEEN ZOEK-EN-VERVANG
De vraagtekst is niet alleen tekst maar ook een sleutel. netto_fotos.js,
netto_bronnen.js en netto_translations_en.js zoeken hun gegevens op met de
letterlijke vraag als sleutel, en de puzzelbestanden bewaren hem als label. Wie
er eentje vergeet, krijgt geen foutmelding: de foto verdwijnt gewoon, of de
vraag blijft in het Nederlands staan terwijl de rest Engels is.

WAT HET WEL EN NIET DOET
Het verandert alleen de tekst, nooit het antwoord. Daardoor blijft elke som
kloppen en vervalt er geen gespeelde score. Klopt het antwoord zelf niet, dan is
dit het verkeerde gereedschap.

De dagpuzzels in de database staan hier buiten: die rijen leven in Supabase en
worden per stuk bijgewerkt. Het script schrijft daarvoor de SQL uit.
"""

import json
import os
import re
import shutil
import sys
from datetime import date

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
MAPPEN = [os.path.join(WORTEL, 'data'), os.path.join(WORTEL, 'website', 'data')]

# oude tekst -> (nieuwe tekst, nieuwe Engelse tekst, nieuwe bewijszin of None)
HERNOEMINGEN = {
    'Hoeveel jaar duurden de Punische oorlogen samen?': (
        'Hoeveel jaar zaten er tussen het begin en het einde van de Punische oorlogen?',
        'How many years were there between the start and the end of the Punic Wars?',
        'Van 264 tot 146 voor Christus is 118 jaar tussen het begin van de Eerste '
        'en het einde van de Derde Punische Oorlog. De drie oorlogen samen duurden '
        '43 jaar; de rest van die periode was vrede.',
    ),
    'In welk jaar viel Carthago aan de Romeinen tijdens de Derde Punische Oorlog?': (
        'In welk jaar voor Christus viel Carthago in handen van de Romeinen?',
        'In which year BC did Carthage fall to the Romans?',
        None,
    ),
    'In welk jaar viel de Berlijnse Muur niet maar werd Japan aangevallen door de VS in Pearl Harbor?': (
        'In welk jaar viel Japan de Amerikaanse vlootbasis Pearl Harbor aan?',
        'In which year did Japan attack the American naval base at Pearl Harbor?',
        'De aanval op Pearl Harbor was een verrassingsaanval door de Japanse Keizerlijke '
        'Marine op de Amerikaanse vlootbasis op 7 december 1941.',
    ),
    'Hoeveel electorale stemmen heeft Californië (2024)?': (
        'Hoeveel kiesmannen heeft Californië (2024)?',
        'How many electoral votes does California have (2024)?',
        None,
    ),
    'Hoeveel electorale stemmen heeft Texas (2024)?': (
        'Hoeveel kiesmannen heeft Texas (2024)?',
        'How many electoral votes does Texas have (2024)?',
        None,
    ),
    'Hoeveel electorale stemmen zijn er nodig om de VS-presidentsverkiezing te winnen?': (
        'Hoeveel kiesmannen zijn er nodig om de Amerikaanse presidentsverkiezing te winnen?',
        'How many electoral votes are needed to win the American presidential election?',
        None,
    ),
    'Hoeveel deelstaten (Bundesländer) telt Duitsland?': (
        'Hoeveel deelstaten telt Duitsland?',
        'How many federal states does Germany have?',
        'De Bondsrepubliek Duitsland is een federatie van zestien deelstaten, in het '
        'Duits Bundesländer of Länder (enkelvoud Land) geheten.',
    ),
}


def vervang_in_bestand(pad, oud, nieuw):
    """Vervang de tekst in een JS-databestand. De vraag staat daar als JSON,
    dus vervangen gebeurt op de JSON-vorm: dan blijven aanhalingstekens en
    accenten precies zoals ze in het bestand staan."""
    with open(pad, encoding='utf-8') as f:
        inhoud = f.read()
    oud_json = json.dumps(oud, ensure_ascii=False)[1:-1]
    nieuw_json = json.dumps(nieuw, ensure_ascii=False)[1:-1]
    aantal = inhoud.count(oud_json)
    if not aantal:
        return 0
    with open(pad, 'w', encoding='utf-8') as f:
        f.write(inhoud.replace(oud_json, nieuw_json))
    return aantal


def main():
    reservekopie = os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_hernoem_{date.today():%Y-%m-%d}_{os.path.basename(REVIEW)}')
    shutil.copy2(REVIEW, reservekopie)

    # Met openpyxl in plaats van pandas, want pandas schrijft een blad opnieuw
    # en gooit daarmee de opmaak, het filter en de bevroren kop weg.
    boek = openpyxl.load_workbook(REVIEW)
    blad = boek['Vragen']
    kop = [c.value for c in blad[1]]
    kolom = {naam: n for n, naam in enumerate(kop, start=1)}
    sql = []

    for oud, (nieuw, engels, bewijs) in HERNOEMINGEN.items():
        # De SQL komt er altijd, ook voor een hernoeming die hier al gedaan is:
        # de database loopt een eigen ronde en kan de oude tekst nog hebben.
        # De kolommen heten in de database question_1 tot en met question_3.
        # q1_label is de naam die de frontend er pas in mapDbDaily aan geeft, en
        # daar zat een eerdere versie van dit script naast: de SQL viel om met
        # 'column "q1_label" does not exist'.
        sql.append('\n'.join(
            "update public.puzzles set question_{i} = {n} where question_{i} = {o};".format(
                i=i, n="'" + nieuw.replace("'", "''") + "'",
                o="'" + oud.replace("'", "''") + "'")
            for i in (1, 2, 3)))

        rijen = [r for r in range(2, blad.max_row + 1)
                 if blad.cell(r, kolom['Vraag NL']).value == oud]
        if len(rijen) != 1:
            print(f'al gedaan: "{oud[:60]}..."')
            continue
        rij = rijen[0]
        blad.cell(rij, kolom['Vraag NL']).value = nieuw
        if bewijs:
            blad.cell(rij, kolom['Bewijszin']).value = bewijs
        print(f'Nr {blad.cell(rij, kolom["Nr"]).value}: {oud}\n      -> {nieuw}')

        for map_ in MAPPEN:
            for naam in sorted(os.listdir(map_)):
                if not naam.endswith('.js'):
                    continue
                pad = os.path.join(map_, naam)
                n = vervang_in_bestand(pad, oud, nieuw)
                if n and map_ == MAPPEN[0]:
                    print(f'      {naam}: {n}x')
        # De vertaling hangt aan de oude sleutel; die moet mee, anders staat de
        # vraag in het Engelse spel opeens in het Nederlands.
        for map_ in MAPPEN:
            pad = os.path.join(map_, 'netto_translations_en.js')
            with open(pad, encoding='utf-8') as f:
                tekst = f.read()
            sleutel = json.dumps(nieuw, ensure_ascii=False)
            patroon = re.compile(re.escape(sleutel) + r':\s*(".*?[^\\]")', re.S)
            treffer = patroon.search(tekst)
            if treffer:
                tekst = tekst[:treffer.start(1)] + json.dumps(engels, ensure_ascii=False) + tekst[treffer.end(1):]
                with open(pad, 'w', encoding='utf-8') as f:
                    f.write(tekst)
        print(f'      vertaling -> {engels}')

    boek.save(REVIEW)

    pad_sql = os.path.join(WORTEL, 'supabase', 'hernoem_vragen.sql')
    with open(pad_sql, 'w', encoding='utf-8') as f:
        f.write('-- Draai dit in de Supabase SQL Editor: de dagpuzzels bewaren de\n'
                '-- vraagtekst in hun eigen rij, dus zonder deze update blijft de\n'
                '-- oude formulering staan in elke daily die al is ingepland.\n\n'
                + '\n\n'.join(sql) + '\n')
    print(f'\nReservekopie: {os.path.basename(reservekopie)}')
    print(f'SQL voor de database: {os.path.relpath(pad_sql, WORTEL)}')


if __name__ == '__main__':
    main()

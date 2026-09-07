# -*- coding: utf-8 -*-
"""Netto — verwijder de onbruikbare reservevragen uit bank, reviewblad en races.

Draaien:  python tools/schrap_vragen.py

WAAROM DIT MEER IS DAN RIJEN WISSEN
De kolom "In gebruik" in het reviewblad zei dat deze 52 vragen nergens gebruikt
werden. Dat klopte niet: die kolom keek alleen naar netto_frontend_puzzles.js en
sloeg de race-sets over. In werkelijkheid zaten ze in 80 speelbare race-puzzels.
Alleen de bank opschonen had die 80 gewoon laten staan, en dan waren vragen die
we onbruikbaar noemen tóch in het spel gebleven.

WAT ER GEBEURT
  bank + reviewblad   de 52 rijen verdwijnen
  race-sets           de 80 puzzels die zo'n vraag bevatten verdwijnen
  race-sets           56 puzzels krijgen de opgeschoonde vraagtekst; hun som
                      blijft kloppen omdat het antwoord niet verandert

WAT ER NIET GEBEURT
109 race-puzzels bevatten een vraag waarvan het antwoord is aangescherpt. Bij die
puzzels zou de som breken: staat er "190 × b = c" en wordt de Cairo Tower 187,
dan klopt de uitkomst niet meer. Die blijven ongemoeid tot ze opnieuw gegenereerd
worden uit de opgeschoonde bank.
"""

import io
import json
import os
import shutil
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from afrondvragen_correcties import CORRECTIES

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK = os.path.join(WORTEL, 'vragen', '1000+ vragen netjes gecategoriseerd.xlsx')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
BACKUP_REVIEW = os.path.join(
    WORTEL, 'vragen', f'_backup_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx')
RACES = os.path.join(WORTEL, 'data', 'netto_race_sets.js')


def main():
    schrap_nrs = {nr for nr, v in CORRECTIES.items() if v[0] == 'schrappen'}

    # De oude vraagteksten staan in de backup; het reviewblad is al herschreven.
    oud = pd.read_excel(BACKUP_REVIEW, sheet_name='Vragen')
    oud_tekst = {int(r['Nr']): str(r['Vraag NL']).strip() for _, r in oud.iterrows()}
    oud_antw = {int(r['Nr']): r['Antwoord'] for _, r in oud.iterrows()}

    weg_teksten = {oud_tekst[nr] for nr in schrap_nrs if nr in oud_tekst}
    # Alleen tekst vervangen waar het antwoord gelijk blijft; anders breekt de som.
    hernoem = {}
    for nr, (status, nv, na, _) in CORRECTIES.items():
        if status == 'schrappen' or not nv or nr not in oud_tekst:
            continue
        if na is not None and str(na) != str(oud_antw.get(nr)):
            continue
        hernoem[oud_tekst[nr]] = nv

    # ---------- reviewblad ----------
    boek = load_workbook(REVIEW)
    ws = boek['Vragen']
    for rij in range(ws.max_row, 1, -1):
        if ws.cell(row=rij, column=1).value in schrap_nrs:
            ws.delete_rows(rij)
    boek.save(REVIEW)
    over = pd.read_excel(REVIEW, sheet_name='Vragen')
    print(f'reviewblad : {len(oud)} -> {len(over)} vragen')

    # ---------- bank ----------
    shutil.copy2(BANK, BANK.replace('.xlsx', f'_voor_schrappen_{date.today():%Y-%m-%d}.xlsx'))
    bank = pd.read_excel(BANK)
    voor = len(bank)
    bank = bank[~bank['Vraag NL'].astype(str).str.strip().isin(weg_teksten)]
    bank.to_excel(BANK, index=False)
    print(f'bank       : {voor} -> {len(bank)} vragen')

    # ---------- race-sets ----------
    ruw = io.open(RACES, encoding='utf-8').read()
    kop, _, rest = ruw.partition('=')
    data = json.loads(rest.strip().rstrip(';'))

    verwijderd, hernoemd = 0, 0
    for naam, groep in data.items():
        houden = []
        for p in groep:
            labels = [str(p.get(sl) or '').strip() for sl in ('q1_label', 'q2_label', 'q3_label')]
            if any(t in weg_teksten for t in labels):
                verwijderd += 1
                continue
            for sl, t in zip(('q1_label', 'q2_label', 'q3_label'), labels):
                if t in hernoem:
                    p[sl] = hernoem[t]
                    hernoemd += 1
            houden.append(p)
        data[naam] = houden

    io.open(RACES, 'w', encoding='utf-8').write(
        kop + '= ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'race-sets  : {verwijderd} puzzels verwijderd, {hernoemd} vraagteksten opgeschoond')
    print('\nomvang per set na opschonen:')
    for naam, groep in sorted(data.items(), key=lambda x: -len(x[1])):
        print(f'  {naam:18s} {len(groep):4d}')


if __name__ == '__main__':
    main()

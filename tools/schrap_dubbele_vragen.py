# -*- coding: utf-8 -*-
"""Netto — verwijder vragen die twee keer in de bank staan.

Draaien:  python tools/schrap_dubbele_vragen.py [--schrijf]

Twee vragen staan er dubbel in, met hetzelfde antwoord: het Suezkanaal en de
Shanghai Tower. Dat is niet alleen slordig maar ook schadelijk, want een puzzel
kan dan twee keer dezelfde vraag krijgen zonder dat de generator dat merkt —
hij vergelijkt op nummer, en de nummers verschillen.

De eerste van elk paar blijft staan, de tweede verdwijnt. Welke van de twee
maakt niet uit; alle andere velden zijn gelijk.
"""

import os
import shutil
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK = os.path.join(WORTEL, 'vragen', '1000+ vragen netjes gecategoriseerd.xlsx')
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')


def main():
    schrijf = '--schrijf' in sys.argv
    r = pd.read_excel(REVIEW, sheet_name='Vragen')
    dub = r[r.duplicated('Vraag NL', keep='first')]
    print(f'dubbel: {len(dub)}')
    for _, x in dub.iterrows():
        print(f"  weg: nr {int(x['Nr'])} — {str(x['Vraag NL'])[:60]} -> {x['Antwoord']}")
    weg_nrs = set(int(n) for n in dub['Nr'])
    weg_tekst = set(str(v) for v in dub['Vraag NL'])

    if not schrijf:
        print('\n(proefdraai — voeg --schrijf toe om op te slaan)')
        return

    # Het reviewblad heeft een Nr-kolom, dus daar kan de tweede rij precies
    # worden aangewezen. De bank heeft die niet; daar wordt van elk paar de
    # tweede voorkomende regel verwijderd.
    for pad, op_nummer in ((REVIEW, True), (BANK, False)):
        shutil.copy2(pad, os.path.join(
            os.path.dirname(pad),
            f'_backup_dubbel_{date.today():%Y-%m-%d}_{os.path.basename(pad)}'))
        boek = load_workbook(pad)
        for blad in boek.worksheets:
            kop = [c.value for c in blad[1]]
            if 'Vraag NL' not in kop:
                continue
            i_v = kop.index('Vraag NL') + 1
            teweg, gezien = [], set()
            for rij in range(2, blad.max_row + 1):
                tekst = blad.cell(row=rij, column=i_v).value
                if op_nummer:
                    if blad.cell(row=rij, column=1).value in weg_nrs:
                        teweg.append(rij)
                elif tekst in weg_tekst:
                    if tekst in gezien:
                        teweg.append(rij)
                    gezien.add(tekst)
            for rij in sorted(teweg, reverse=True):
                blad.delete_rows(rij)
            if teweg:
                print(f'{os.path.basename(pad)} / {blad.title}: {len(teweg)} weg')
        boek.save(pad)

    na = pd.read_excel(REVIEW, sheet_name='Vragen')
    print(f'\nvragen over: {len(na)} | unieke teksten: {na["Vraag NL"].nunique()}')


if __name__ == '__main__':
    main()

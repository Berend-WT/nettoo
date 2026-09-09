"""Maak bronvermeldingen beschikbaar zonder de vragen of antwoorden te wijzigen.

Draaien: python tools/maak_bronnenbestand.py
Leest alleen het reviewblad. Exacte frontendlabels met een vraagnummer krijgen
een extra sleutel; inhoudelijk afwijkende teksten worden niet gegokt.
"""

import argparse
import json
import re
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

WORTEL = Path(__file__).resolve().parents[1]


def maak_bronnen(werkmap, puzzelbestand):
    boek = load_workbook(werkmap, read_only=True, data_only=True)
    bronnen = {}
    aantal = 0
    met_bron = 0
    try:
        rijen = boek['Vragen'].iter_rows(values_only=True)
        kop = next(rijen)
        kolommen = [kop.index(naam) for naam in ('Vraag NL', 'Bron (geverifieerd)', 'Bewijszin')]
        for rij in rijen:
            vraag, bron, uitleg = (rij[i] for i in kolommen)
            if not vraag:
                continue
            aantal += 1
            bron = str(bron or '').strip()
            uitleg = str(uitleg or '').strip()
            try:
                url = urlsplit(bron)
                geldig = url.scheme in ('http', 'https') and bool(url.hostname) and not any(c.isspace() for c in bron)
            except ValueError:
                geldig = False
            # Een plaatshouder is geen bewijs. Hij stond wel in het paneel "Waar
            # komt dit vandaan?", waar de speler dan las dat het antwoord ergens
            # op de pagina staat. Liever geen paneel dan een lege belofte.
            if 'bestaande bron opgehaald' in uitleg.casefold():
                continue
            if not geldig or not uitleg:
                continue
            vermelding = {'bron': bron, 'uitleg': uitleg}
            if vraag in bronnen and bronnen[vraag] != vermelding:
                raise ValueError(f'Tegenstrijdige bronnen bij vraag: {vraag}')
            bronnen[vraag] = vermelding
            met_bron += 1
    finally:
        boek.close()
    unieke_bronnen = len(bronnen)
    tekst = Path(puzzelbestand).read_text(encoding='utf-8-sig')
    puzzels = json.loads(tekst.split('=', 1)[1].strip().rstrip(';'))
    labels = {p[f'q{i}_label'] for reeks in puzzels.values() if isinstance(reeks, list) for p in reeks for i in (1, 2, 3)}
    for label in sorted(labels):
        zonder_nummer = re.sub(r'^\d+\.\s+', '', label)
        if label not in bronnen and zonder_nummer in bronnen:
            bronnen[label] = bronnen[zonder_nummer]
    return bronnen, {'vragen': aantal, 'met_bron': met_bron, 'unieke_bronnen': unieke_bronnen, 'frontendlabels': len(labels), 'frontend_met_bron': sum(label in bronnen for label in labels)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--werkmap', type=Path, default=WORTEL / 'vragen/vragen_review_compleet.xlsx')
    parser.add_argument('--puzzels', type=Path, default=WORTEL / 'data/netto_frontend_puzzles.js')
    parser.add_argument('--uitvoer', type=Path, default=WORTEL / 'data/netto_bronnen.js')
    args = parser.parse_args()
    bronnen, telling = maak_bronnen(args.werkmap, args.puzzels)
    args.uitvoer.parent.mkdir(parents=True, exist_ok=True)
    inhoud = '// Gegenereerd uit het reviewblad; wijzig bronnen in de werkmap.\nwindow.NETTO_BRONNEN = '
    inhoud += json.dumps(bronnen, ensure_ascii=False, sort_keys=True, indent=2) + ';\n'
    args.uitvoer.write_text(inhoud, encoding='utf-8')
    # De bestaande synchronisatietool kent dit nieuwe bestand nog niet.
    # Houd de deploykopie daarom hier gelijk zonder die tool te wijzigen.
    if args.uitvoer.resolve() == (WORTEL / 'data/netto_bronnen.js').resolve():
        spiegel = WORTEL / 'website/data/netto_bronnen.js'
        spiegel.parent.mkdir(parents=True, exist_ok=True)
        spiegel.write_text(inhoud, encoding='utf-8')
    print(json.dumps(telling, ensure_ascii=False))


if __name__ == '__main__':
    main()

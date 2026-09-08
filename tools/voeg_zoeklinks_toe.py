# -*- coding: utf-8 -*-
"""Netto — zet klikbare zoeklinks in de foto-invullijst.

Draaien:  python tools/voeg_zoeklinks_toe.py

Zoeken naar een bruikbare foto kost de meeste tijd, niet het plakken. Deze
lijst neemt dat werk weg met twee links per vraag.

  Zoek op Commons   Gaat rechtstreeks naar Wikimedia Commons, gefilterd op
                    afbeeldingen. Alles wat je daar ziet is herbruikbaar en
                    heeft een controleerbare licentie.

  Foto in de bron   Alleen bij vragen waarvan de bron een Wikipedia-artikel is,
                    en dat zijn de meeste. Het artikel gaat per definitie over
                    het juiste onderwerp, en de afbeeldingen erin staan vrijwel
                    altijd op Commons. Dit is de snelste weg: artikel openen,
                    op de foto klikken, doorklikken naar de bestandspagina.

De zoekterm komt uit de Engelse zoekhulp als die er is, anders uit de
Nederlandse vraag. Commons is grotendeels Engelstalig, dus een Engelse term
levert meer op.
"""

import os
import re
import shutil
import urllib.parse

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DOEL = os.path.join(WORTEL, 'vragen', 'vragen_bronnen_fotos.xlsx')
ONEDRIVE = os.path.join(os.path.expanduser('~'), 'OneDrive - Driestar-Wartburg')

VOLGORDE = {'daily': 0, 'puzzel': 1}

# Woorden die de vraag opbouwen maar niets over het onderwerp zeggen. Ze eruit
# halen scheelt ruis in de zoekopdracht.
RUIS = {
    'hoeveel', 'how', 'many', 'much', 'what', 'welk', 'welke', 'hoe', 'wat',
    'een', 'de', 'het', 'van', 'in', 'op', 'is', 'are', 'the', 'a', 'an', 'of',
    'zijn', 'er', 'en', 'and', 'or', 'of', 'die', 'dat', 'met', 'with', 'voor',
    'per', 'bij', 'aan', 'te', 'tot', 'als', 'does', 'do', 'has', 'have',
    'werd', 'wordt', 'heeft', 'hebben', 'telt', 'staan', 'staat', 'duurt',
    'ongeveer', 'about', 'approximately', 'gemiddeld', 'average', 'standaard',
    'standard', 'volgens', 'according', 'totaal', 'total', 'jaar', 'year',
    'meter', 'metres', 'meters', 'kilometer', 'kilometers', 'centimeter',
    'millimeter', 'kilogram', 'gram', 'liter', 'litres', 'ton', 'tons',
    'procent', 'percent', 'percentage', 'graden', 'degrees', 'seconde',
    'seconden', 'seconds', 'minuut', 'minuten', 'minutes', 'uur', 'hours',
    'dagen', 'dag', 'days', 'day', 'maanden', 'months', 'miljoen', 'million',
    'miljard', 'billion', 'duizend', 'thousand', 'aantal', 'number', 'lang',
    'long', 'hoog', 'high', 'tall', 'diep', 'deep', 'breed', 'wide', 'groot',
    'large', 'zwaar', 'heavy', 'maximaal', 'wereldwijd', 'worldwide', 'eerste',
    'first', 'stand', 'klassiek', 'classic',
}


def zoekterm(nl, en):
    bron = en if isinstance(en, str) and en.strip() else nl
    woorden = re.findall(r"[A-Za-zÀ-ÿ'’-]{3,}", str(bron))
    kern = [w for w in woorden if w.lower() not in RUIS]
    return ' '.join(kern[:5]) or str(nl)[:60]


def commons_url(term):
    q = urllib.parse.quote_plus(term)
    # search=...&title=Special:MediaSearch&type=image geeft direct de
    # afbeeldingenweergave in plaats van een tekstuele trefferlijst.
    return (f'https://commons.wikimedia.org/w/index.php?search={q}'
            f'&title=Special:MediaSearch&type=image')


def main():
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    en = d['Vraag EN (zoekhulp)'] if 'Vraag EN (zoekhulp)' in d.columns else pd.Series(
        [None] * len(d), index=d.index)

    termen = [zoekterm(nl, e) for nl, e in zip(d['Vraag NL'], en)]
    bron = d['Bron (geverifieerd)'].fillna('')
    is_wiki = bron.str.contains('wikipedia.org/wiki/', regex=False)

    uit = pd.DataFrame({
        'Nr': d['Nr'],
        'In gebruik': d['In gebruik'],
        'Categorie': d['Categorie'],
        'Vraag NL': d['Vraag NL'],
        'Antwoord': d['Antwoord'],
        'Zoek op Commons': [commons_url(t) for t in termen],
        'Foto in de bron': [b if w else '' for b, w in zip(bron, is_wiki)],
        'Fotolink (Commons)': None,
        'Bron': bron,
        'Waarom die bron': d['Bewijszin'],
    })
    uit['_r'] = uit['In gebruik'].map(VOLGORDE).fillna(2)
    uit = uit.sort_values(['_r', 'Categorie', 'Nr']).drop(columns=['_r'])
    uit.to_excel(DOEL, index=False, sheet_name='Fotos')

    wb = load_workbook(DOEL)
    ws = wb['Fotos']
    br = {'Nr': 6, 'In gebruik': 11, 'Categorie': 22, 'Vraag NL': 58, 'Antwoord': 12,
          'Zoek op Commons': 18, 'Foto in de bron': 18, 'Fotolink (Commons)': 46,
          'Bron': 44, 'Waarom die bron': 60}
    for i, k in enumerate(uit.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = br.get(k, 18)
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid',
                             fgColor='2E7D32' if k == 'Fotolink (Commons)' else '1F3864')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30

    kolommen = list(uit.columns)
    i_zoek = kolommen.index('Zoek op Commons') + 1
    i_art = kolommen.index('Foto in de bron') + 1
    i_vul = kolommen.index('Fotolink (Commons)') + 1
    i_geb = kolommen.index('In gebruik') + 1

    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(kolommen) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=kolommen[kol - 1] in
                                    ('Vraag NL', 'Waarom die bron'))
        # De lange adressen zijn onleesbaar in een cel; als klikbare tekst
        # blijft de kolom smal en is de bedoeling meteen duidelijk.
        for kol, tekst in ((i_zoek, 'zoek foto'), (i_art, 'open artikel')):
            c = ws.cell(row=rij, column=kol)
            if c.value:
                c.hyperlink = c.value
                c.value = tekst
                c.font = Font(name='Arial', size=10, color='0563C1', underline='single')
        ws.cell(row=rij, column=i_vul).fill = PatternFill('solid', fgColor='FFF9E0')
        g = ws.cell(row=rij, column=i_geb).value
        if g in VOLGORDE:
            ws.cell(row=rij, column=i_geb).fill = PatternFill(
                'solid', fgColor='D6EAD6' if g == 'daily' else 'E8F0DC')

    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(kolommen))}{ws.max_row}'
    wb.save(DOEL)

    if os.path.isdir(ONEDRIVE):
        shutil.copy2(DOEL, os.path.join(ONEDRIVE, os.path.basename(DOEL)))
    print(f'{len(uit)} vragen')
    print(f'met een Wikipedia-artikel als bron: {int(is_wiki.sum())}')
    print(f'-> {DOEL}')


if __name__ == '__main__':
    main()

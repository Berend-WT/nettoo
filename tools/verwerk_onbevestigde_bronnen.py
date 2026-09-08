# -*- coding: utf-8 -*-
"""Netto — zet de herbeoordeelde bronnen door naar het reviewblad.

Draaien:  python tools/verwerk_onbevestigde_bronnen.py

De vijf blokken samen dekken de 316 vragen die wel een bron hadden, maar waarvan
die bron het antwoord niet bevestigde. Zelfde oordelen als bij de reservevragen:

  bevestigd    Bron plus toelichting; geldt als geverifieerd.
  natellen     Bron klopt maar noemt het getal niet als getal.
  veroudert    Bron klopt, maar het antwoord verschuift met de tijd.
  fout         Het antwoord klopt niet en moet eerst worden gecorrigeerd.
  onbruikbaar  De vraag zelf deugt niet.

Antwoorden worden niet aangepast; dat breekt bij vragen in een puzzel de som.
Naast het reviewblad komt er een apart werkblad met alleen deze 316, gesorteerd
op oordeel, zodat de fouten vooraan staan.
"""

import os
import shutil
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bronnen_onbevestigd1 import O
from bronnen_onbevestigd2 import O2
from bronnen_onbevestigd3 import O3
from bronnen_onbevestigd4 import O4
from bronnen_onbevestigd5 import O5

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
DOEL = os.path.join(WORTEL, 'vragen', 'herbeoordeelde_bronnen.xlsx')

ALLES = {}
for blok in (O, O2, O3, O4, O5):
    ALLES.update(blok)

KLEUR = {'fout': 'F8CBCB', 'onbruikbaar': 'F2D5D5', 'veroudert': 'FBE9A5',
         'natellen': 'DDE7F5', 'bevestigd': 'D6EAD6'}
RANG = ['fout', 'onbruikbaar', 'veroudert', 'natellen', 'bevestigd']
VERTROUWEN = {'bevestigd': 'handmatig', 'natellen': 'natellen',
              'veroudert': 'veroudert', 'fout': 'antwoord fout'}


def main():
    d = pd.read_excel(REVIEW, sheet_name='Vragen')
    verwacht = set(d.loc[d['Bron (geverifieerd)'].isna() & d['Bron (bestaand)'].notna(), 'Nr'])
    gemist = verwacht - set(ALLES)
    extra = set(ALLES) - verwacht
    if gemist:
        print(f'LET OP — niet beoordeeld: {sorted(gemist)}')
    if extra:
        print(f'LET OP — buiten de selectie: {sorted(extra)}')

    rijen = []
    for _, r in d.iterrows():
        nr = int(r['Nr'])
        if nr not in ALLES:
            continue
        oordeel, bron, toelichting = ALLES[nr]
        rijen.append({'Nr': nr, 'Oordeel': oordeel, 'In gebruik': r['In gebruik'],
                      'Categorie': r['Categorie'], 'Vraag NL': r['Vraag NL'],
                      'Antwoord': r['Antwoord'], 'Bron (oud)': r['Bron (bestaand)'],
                      'Bron (nieuw)': bron, 'Toelichting': toelichting,
                      'Jouw besluit': None})

    uit = pd.DataFrame(rijen)
    uit['_r'] = uit['Oordeel'].apply(RANG.index)
    uit = uit.sort_values(['_r', 'Nr']).drop(columns=['_r'])
    uit.to_excel(DOEL, index=False, sheet_name='Herbeoordeeld')

    wb = load_workbook(DOEL)
    ws = wb['Herbeoordeeld']
    br = {'Nr': 6, 'Oordeel': 14, 'In gebruik': 11, 'Categorie': 22, 'Vraag NL': 56,
          'Antwoord': 12, 'Bron (oud)': 44, 'Bron (nieuw)': 50, 'Toelichting': 74,
          'Jouw besluit': 18}
    for i, k in enumerate(uit.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = br.get(k, 18)
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    i_o = list(uit.columns).index('Oordeel') + 1
    for rij in range(2, ws.max_row + 1):
        for kol in range(1, len(uit.columns) + 1):
            c = ws.cell(row=rij, column=kol)
            c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=uit.columns[kol - 1] in ('Vraag NL', 'Toelichting'))
        o = ws.cell(row=rij, column=i_o).value
        ws.cell(row=rij, column=i_o).fill = PatternFill('solid', fgColor=KLEUR.get(o, 'FFFFFF'))
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(uit.columns))}{ws.max_row}'
    wb.save(DOEL)

    shutil.copy2(REVIEW, os.path.join(
        os.path.dirname(REVIEW),
        f'_backup_herbeoordeeld_{date.today():%Y-%m-%d}_vragen_review_compleet.xlsx'))
    boek = load_workbook(REVIEW)
    blad = boek['Vragen']
    kol = [c.value for c in blad[1]]
    i_bron = kol.index('Bron (geverifieerd)') + 1
    i_bew = kol.index('Bewijszin') + 1
    i_vert = kol.index('Vertrouwen bron') + 1
    i_st = kol.index('Status oude bron') + 1

    n = 0
    for rij in range(2, blad.max_row + 1):
        nr = blad.cell(row=rij, column=1).value
        if nr not in ALLES:
            continue
        oordeel, bron, toelichting = ALLES[nr]
        if not bron:
            blad.cell(row=rij, column=i_st).value = f'geen bron mogelijk ({oordeel})'
            blad.cell(row=rij, column=i_bew).value = toelichting[:400]
            continue
        blad.cell(row=rij, column=i_bron).value = bron
        blad.cell(row=rij, column=i_bew).value = toelichting[:400]
        blad.cell(row=rij, column=i_vert).value = VERTROUWEN.get(oordeel, oordeel)
        if oordeel == 'fout':
            blad.cell(row=rij, column=i_st).value = 'antwoord fout, zie toelichting'
        elif oordeel == 'onbruikbaar':
            blad.cell(row=rij, column=i_st).value = 'vraag deugt niet, zie toelichting'
        else:
            blad.cell(row=rij, column=i_st).value = 'bron vervangen'
        n += 1
    boek.save(REVIEW)

    print(uit['Oordeel'].value_counts().to_string())
    print(f'\nbronnen weggeschreven: {n}')

    na = pd.read_excel(REVIEW, sheet_name='Vragen')
    ver = na['Bron (geverifieerd)'].notna()
    print(f'kolom "Bron (geverifieerd)" gevuld: {int(ver.sum())} van {len(na)}')
    print(f'-> {DOEL}')


if __name__ == '__main__':
    main()

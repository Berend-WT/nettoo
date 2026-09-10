# -*- coding: utf-8 -*-
"""Netto — zoek de hoofdafbeelding bij het ONDERWERP van de vraag.

Draaien:  python tools/zoek_onderwerpafbeeldingen.py [--alleen-rood] [--aantal N]

WAAROM DIT ER NOG BIJ MOET
zoek_hoofdafbeeldingen.py pakt de infoboxfoto van het bronartikel. Dat werkt
alleen als de bron een Wikipedia-artikel is. Staat het antwoord op
coca-cola.com, who.int of esbnyc.com, dan is er geen infobox en valt de vraag
terug op de oude tekstzoekronde. Die zoekt op Commons, waar de licentiefilter
bijna alles wegstreept behalve publiek domein, en publiek domein is daar oud
drukwerk. Zo kreeg de vraag over een blikje cola drie negentiende-eeuwse
advertenties als keuze.

De oplossing is niet beter zoeken maar beter kijken: het onderwerp van de vraag
heeft zelf bijna altijd een Wikipedia-artikel, en dat artikel heeft de foto die
je verwacht. "Coca-Cola" levert een blikje, "Empire State Building" het gebouw.

HOE HET ONDERWERP WORDT GERADEN
Eigennamen in de vraag, van lang naar kort, plus als terugval het zelfstandig
naamwoord dat geteld wordt. De eerste die op Wikipedia een artikel met een
infoboxfoto oplevert, wint. Wat niets oplevert blijft leeg; een foute foto is
erger dan geen foto.
"""

import argparse
import json
import os
import re
import sys
import time
import urllib.parse

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import zoek_hoofdafbeeldingen as zh

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
FOTOS = os.path.join(WORTEL, 'fotos')
HOOFD = os.path.join(FOTOS, 'hoofdafbeeldingen.json')
UIT = os.path.join(FOTOS, 'onderwerpafbeeldingen.json')
PAUZE = 0.6  # Wikipedia geeft boven de tien verzoeken per seconde zwijgend niets terug

# Woorden die met een hoofdletter beginnen maar geen onderwerp zijn.
GEEN_NAAM = {'Hoeveel', 'Welke', 'Welk', 'Hoe', 'Wat', 'In', 'Er', 'De', 'Het', 'Een'}
# Woorden die wel achteraan de vraag staan maar nooit het onderwerp zijn.
# Deze lijst is niet bedacht maar afgeleid: Berend keurde 34 voorstellen en
# vrijwel alles wat hij wegstreepte was gezocht op een maateenheid, een
# bijwoord of een werkwoord. "Hoe hoog was de hoogste duik in meters" leverde
# een foto van een duimstok, "hoeveel legt een kip" een luchtfoto van Getafe.
STAART = {
    # maten
    'meter', 'meters', 'centimeter', 'centimeters', 'millimeter', 'millimeters',
    'kilometer', 'kilometers', 'kilo', 'kilogram', 'gram', 'ton', 'tonnen',
    'liter', 'liters', 'graden', 'hertz', 'procent', 'seconde', 'seconden',
    'minuut', 'minuten', 'uur', 'uren', 'dag', 'dagen', 'week', 'weken',
    'maand', 'maanden', 'jaar', 'jaren', 'eeuw', 'eeuwen', 'euro', 'dollar',
    # bijwoorden en overtreffende trappen
    'ooit', 'exact', 'precies', 'gemiddeld', 'gemiddelde', 'maximaal', 'minimaal',
    'snelst', 'hoogst', 'langst', 'grootst', 'zwaarst', 'diepst', 'kleinst',
    'meest', 'totaal', 'samen', 'ongeveer', 'wereldwijd', 'jaarlijks',
    'dagelijks', 'tegelijk', 'momenteel', 'huidige', 'benadering', 'afgerond',
    'stand', 'ruwweg', 'naar', 'schatting', 'volgens', 'begin', 'start',
    # werkwoordsvormen
    'gelanceerd', 'gepubliceerd', 'opgenomen', 'uitgestrekt', 'geproduceerd',
    'gebouwd', 'gemaakt', 'verkocht', 'gewonnen', 'gespeeld', 'geschreven',
    'gevonden', 'gebruikt', 'geteld', 'bedekt', 'legt', 'meet', 'weegt',
    'telt', 'duurt', 'bevat', 'kost', 'haalt', 'staat', 'ligt', 'loopt',
    'rijdt', 'vliegt', 'zwemt', 'klopt', 'draait', 'heeft', 'hebben',
    # te vaag om een artikel mee te vinden
    'veld', 'land', 'leven', 'ding', 'dingen', 'deel', 'delen', 'aantal',
    'soort', 'soorten', 'stuk', 'stuks', 'keer', 'plek', 'systeem', 'wereld',
    'volledig', 'volledige', 'compleet', 'complete', 'hele', 'gehele',
}
TELWOORD_OVERSLAAN = {'verschillende', 'officiele', 'officiële', 'individuele', 'erkende',
    'bekende', 'gepubliceerde', 'complete', 'standaard', 'totale', 'unieke', 'actieve',
    'echte', 'grote', 'kleine', 'afzonderlijke', 'belangrijkste', 'centrale', 'natuurlijke',
    'zware', 'gewone', 'huidige', 'oorspronkelijke', 'duizend', 'miljoen', 'miljard',
    'procent', 'gram', 'kilo', 'kilogram', 'meter', 'kilometer', 'liter', 'ton', 'jaar'}


def onderwerpen(vraag):
    """Kandidaat-onderwerpen, van meest naar minst kansrijk."""
    uit = []
    # Aaneengesloten reeksen woorden met een hoofdletter: "Empire State Building",
    # "Coca-Cola". Het eerste woord van de zin telt niet mee als naam.
    woorden = re.findall(r"[\w'’-]+", vraag, flags=re.UNICODE)
    reeks = []
    for n, woord in enumerate(woorden):
        naam = woord[:1].isupper() and woord not in GEEN_NAAM and not (n == 0)
        if naam:
            reeks.append((n, woord))
        else:
            if reeks:
                uit.append((reeks[0][0], ' '.join(w for _, w in reeks)))
            reeks = []
    if reeks:
        uit.append((reeks[0][0], ' '.join(w for _, w in reeks)))
    # Langste naam eerst, en bij gelijke lengte de laatste in de zin. Bij "de
    # Slag om Gettysburg" zijn "Slag" en "Gettysburg" allebei een woord, maar
    # het tweede zegt waar het over gaat en het eerste levert een dorp in Chili.
    uit.sort(key=lambda p: (-len(p[1].split()), -p[0]))
    uit = [naam for _, naam in uit]
    # Zonder eigennaam staat het onderwerp meestal achteraan: "Hoeveel kilo
    # weegt een KONINGSPINGUIN", "Hoeveel toetsen heeft een standaard PIANO".
    # Bijwoorden als "ongeveer" en "wereldwijd" hangen er los achter en tellen
    # dus niet mee.
    achteraan = [w for w in re.findall(r'[a-zA-Zà-ÿ]{4,}', vraag.split('?')[0])
                 if w.lower() not in TELWOORD_OVERSLAAN and w.lower() not in STAART]
    if achteraan:
        uit.append(achteraan[-1])
    # Daarna pas het lijdend voorwerp: bij "Hoeveel bulten heeft een dromedaris"
    # gaat het om de dromedaris, niet om de bulten.
    m = re.search(r'\b(?:heeft|hebben|telt|tellen|bevat|kent|zit er in|zitten er in)\s+'
                  r'(?:een |de |het |)([a-zA-Zà-ÿ]{4,})', vraag, re.I)
    if m and m.group(1).lower() not in TELWOORD_OVERSLAAN:
        uit.append(m.group(1))
    # Terugval: het getelde zelfstandig naamwoord ("Hoeveel bulten heeft ...").
    m = re.search(r'\bhoeveel\s+(.+)', vraag, re.I)
    if m:
        for woord in re.findall(r'[a-zA-Zà-ÿ]+', m.group(1)):
            if len(woord) >= 4 and woord.lower() not in TELWOORD_OVERSLAAN:
                uit.append(woord)
                break
    return uit


def artikel(taal, zoekterm):
    """Titel van het beste zoekresultaat, of niets."""
    d = zh.haal(f'{taal}.wikipedia.org', {
        'action': 'query', 'format': 'json', 'formatversion': '2',
        'list': 'search', 'srsearch': zoekterm, 'srlimit': '1', 'srnamespace': '0',
    })
    treffers = (d.get('query', {}) or {}).get('search', []) or []
    return treffers[0]['title'] if treffers else None


def zoek(vraag):
    for term in onderwerpen(vraag):
        for taal in ('nl', 'en'):
            titel = artikel(taal, term)
            time.sleep(PAUZE)
            if not titel:
                continue
            # Doorverwijspagina's hebben geen eigen beeld en leiden altijd fout.
            if re.search(r'\((doorverwijspagina|disambiguation)\)', titel, re.I):
                continue
            url = f'https://{taal}.wikipedia.org/wiki/' + urllib.parse.quote(titel.replace(' ', '_'))
            foto = zh.hoofdafbeelding(url)
            time.sleep(PAUZE)
            if foto:
                foto['gezocht'] = term
                foto['artikel'] = url
                return foto
    return None


def main():
    ontleder = argparse.ArgumentParser()
    ontleder.add_argument('--aantal', type=int, default=0, help='0 = alles')
    ontleder.add_argument('--alleen-rood', action='store_true',
                          help='alleen vragen die nu een foto uit de oude zoekronde tonen')
    args = ontleder.parse_args()

    hoofd = json.load(open(HOOFD, encoding='utf-8')) if os.path.exists(HOOFD) else {}
    gedaan = json.load(open(UIT, encoding='utf-8')) if os.path.exists(UIT) else {}

    blad = openpyxl.load_workbook(REVIEW, read_only=True)['Vragen']
    rijen = list(blad.iter_rows(values_only=True))
    kop = list(rijen[0])
    k = {naam: n for n, naam in enumerate(kop)}

    rood = set()
    if args.alleen_rood:
        keuze = openpyxl.load_workbook(
            os.path.join(WORTEL, 'vragen', 'fotokeuze.xlsx'), read_only=True)['Fotokeuze']
        kr = list(keuze.iter_rows(values_only=True))
        kk = {naam: n for n, naam in enumerate(kr[0])}
        rood = {int(r[kk['Nr']]) for r in kr[1:]
                if r[kk['Nr']] is not None and r[kk['Nu in het spel']] == 'oude zoekronde'}
        print(f'{len(rood)} vragen met een foto uit de oude zoekronde')

    open_nog = []
    for rij in rijen[1:]:
        nr = int(rij[k['Nr']])
        if str(nr) in gedaan:
            continue
        if rood and nr not in rood:
            continue
        if not rood and (hoofd.get(str(nr), {}).get('kandidaten') or []):
            continue  # heeft al een goede hoofdafbeelding uit het bronartikel
        open_nog.append((nr, str(rij[k['Vraag NL']])))
    if args.aantal:
        open_nog = open_nog[:args.aantal]

    print(f'{len(open_nog)} vragen zoeken ...')
    for n, (nr, vraag) in enumerate(open_nog, start=1):
        foto = zoek(vraag)
        gedaan[str(nr)] = {'kandidaten': [foto] if foto else [],
                           'herkomst': 'onderwerpartikel' if foto else 'niets'}
        merk = foto['titel'].replace('File:', '')[:52] if foto else '-'
        print(f'  {n:>3}/{len(open_nog)} Nr {nr:<5} {merk:<54} {vraag[:44]}')
        if n % 10 == 0 or n == len(open_nog):
            with open(UIT, 'w', encoding='utf-8') as f:
                json.dump(gedaan, f, ensure_ascii=False, indent=1)
    raak = sum(1 for v in gedaan.values() if v['kandidaten'])
    print(f'\n{raak} van {len(gedaan)} met een afbeelding -> {os.path.relpath(UIT, WORTEL)}')


if __name__ == '__main__':
    main()

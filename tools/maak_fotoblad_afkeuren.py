# -*- coding: utf-8 -*-
"""Netto — werkblad waarin je alleen de slechte foto's wegstreept.

Draaien:  python tools/maak_fotoblad_afkeuren.py

WAAROM ANDERSOM
Het gewone keuzeblad vraagt per vraag een keuze uit drie kandidaten. Voor de
vragen waarvan de bron een Wikipedia-artikel is hoeft dat niet: die krijgen de
infoboxfoto en die klopt per definitie. Wat overblijft zijn de vragen waar de
foto is geraden uit de vraagtekst. Gemeten op een steekproef van 26 klopt daar
ongeveer driekwart van.

Driekwart is te weinig om ongezien toe te passen, maar te veel om alles opnieuw
te laten uitzoeken. Dus staat er per rij één voorstel, en streep je alleen weg
wat niet deugt — ongeveer een op de vier. Dat is drie keer minder werk dan
kiezen, en wat overblijft is door een mens goedgekeurd in plaats van door een
script geraden.

INVULLEN, in de kolom "Oordeel"
  1                 deze foto is goed
  0                 deze foto deugt niet
  een Commons-URL   gebruik deze in plaats van het voorstel
  vrije tekst       een opmerking; die komt bij het verwerken op het scherm
  leeg              nog niet bekeken, er verandert niets

Een eerdere versie vroeg alleen om weg te strepen, met een apart vakje voor hoe
ver je gekomen was. Berend vulde spontaan 1 en 0 in, en dat is beter: een
ingevulde regel is dan altijd een beslissing en een lege regel altijd "nog niet
bekeken". Wat je beoordeeld hebt komt in een volgende ronde niet terug.
"""

import os
import sys

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# Het downloaden en schalen van miniaturen staat al in het keuzeblad-script;
# tweemaal dezelfde code onderhouden is vragen om verschil.
from maak_fotokeuzeblad import haal_miniatuur

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REVIEW = os.path.join(WORTEL, 'vragen', 'vragen_review_compleet.xlsx')
FOTOS = os.path.join(WORTEL, 'fotos')
MINIATUREN = os.path.join(FOTOS, 'assets', 'kandidaten')
UIT = os.path.join(WORTEL, 'vragen', 'fotos_afkeuren.xlsx')
RANG = {'daily': 0, 'puzzel': 1, 'race': 2, 'breinkraker': 3}
KOLOMMEN = ['Nr', 'In gebruik', 'Vraag', 'Antwoord', 'Foto', 'Oordeel',
            'Gezocht op', 'Naamtreffer', 'Bestand', 'Commons']
BREEDTES = [6, 12, 62, 10, 26, 8, 18, 12, 34, 12]


def main():
    import json
    hoofd = json.load(open(os.path.join(FOTOS, 'hoofdafbeeldingen.json'), encoding='utf-8'))
    onderwerp = json.load(open(os.path.join(FOTOS, 'onderwerpafbeeldingen.json'), encoding='utf-8'))
    met_infobox = {nr for nr, v in hoofd.items() if v.get('kandidaten')}

    # Wat je al beoordeeld hebt komt niet terug. Een oordeel staat of in de
    # kolom Keuze van het keuzeblad, of - bij een zelf opgezocht adres - in
    # handmatige_fotos.json. Zonder deze stap krijg je bij elke ronde dezelfde
    # rijen opnieuw voorgeschoteld.
    beslist = set()
    pad_handmatig = os.path.join(FOTOS, 'handmatige_fotos.json')
    if os.path.exists(pad_handmatig):
        beslist |= set(json.load(open(pad_handmatig, encoding='utf-8')))
    pad_keuze = os.path.join(WORTEL, 'vragen', 'fotokeuze.xlsx')
    if os.path.exists(pad_keuze):
        kb = openpyxl.load_workbook(pad_keuze, read_only=True)['Fotokeuze']
        kr = list(kb.iter_rows(values_only=True))
        kk = {naam: n for n, naam in enumerate(kr[0])}
        beslist |= {str(int(x[kk['Nr']])) for x in kr[1:]
                    if x[kk['Nr']] is not None and x[kk['Keuze']] is not None}
    print(f'{len(beslist)} vragen zijn al beoordeeld en blijven buiten dit blad')

    blad = openpyxl.load_workbook(REVIEW, read_only=True)['Vragen']
    rijen = list(blad.iter_rows(values_only=True))
    k = {naam: n for n, naam in enumerate(rijen[0])}

    te_doen = []
    for r in rijen[1:]:
        if r[k['Nr']] is None:
            continue
        nr = str(int(r[k['Nr']]))
        gebruik = r[k['In gebruik']]
        if gebruik not in RANG or nr in met_infobox or nr in beslist:
            continue
        kandidaat = (onderwerp.get(nr, {}).get('kandidaten') or [None])[0]
        if not kandidaat:
            continue
        # Sorteren op zichtbaarheid, en daarbinnen op de naamtreffer: klopt de
        # naam van het gevonden artikel met de vraag, dan is het voorstel
        # sterker. Zo staat het beste bovenaan en is een half nagelopen blad
        # nog steeds het meest bruikbare deel.
        te_doen.append((RANG[gebruik], 0 if kandidaat.get('naamtreffer') else 1,
                        int(nr), gebruik, str(r[k['Vraag NL']]),
                        r[k['Antwoord']], kandidaat))
    te_doen.sort()
    print(f'{len(te_doen)} voorstellen om na te lopen')

    boek = openpyxl.Workbook()
    ws = boek.active
    ws.title = 'Fotos'
    donker = PatternFill('solid', fgColor='1F3864')
    rood = PatternFill('solid', fgColor='FCE4E4')
    for n, (naam, breedte) in enumerate(zip(KOLOMMEN, BREEDTES), start=1):
        cel = ws.cell(1, n, naam)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = donker
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(n)].width = breedte
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'C2'

    meegenomen = 0
    for rij, (_, _, nr, gebruik, vraag, antwoord, kandidaat) in enumerate(te_doen, start=2):
        waarden = [nr, gebruik, vraag, antwoord, None, None,
                   kandidaat.get('gezocht', ''),
                   'ja' if kandidaat.get('naamtreffer') else 'nee',
                   kandidaat['titel'].replace('File:', ''), 'commons']
        for kol, waarde in enumerate(waarden, start=1):
            cel = ws.cell(rij, kol, waarde)
            cel.alignment = Alignment(vertical='center', wrap_text=(kol == 3))
            if KOLOMMEN[kol - 1] == 'Oordeel':
                cel.fill = rood
            if KOLOMMEN[kol - 1] == 'Commons' and kandidaat.get('pagina'):
                cel.hyperlink = kandidaat['pagina']
                cel.font = Font(color='0563C1', underline='single')
        pad = os.path.join(MINIATUREN, f'{nr}_0.jpg')
        # Het keuzeblad haalt alleen miniaturen op voor daily- en puzzelvragen.
        # Race en breinkrakers staan hier ook in, dus die ontbreken nog. Een rij
        # zonder plaatje valt niet te beoordelen, dus die halen we alsnog.
        if not os.path.exists(pad) and kandidaat.get('miniatuur'):
            haal_miniatuur(kandidaat['miniatuur'], pad)
        if os.path.exists(pad):
            try:
                beeld = XlImage(pad)
                ws.add_image(beeld, f'E{rij}')
                meegenomen += 1
            except Exception:
                pass
        ws.row_dimensions[rij].height = 96

    ws.auto_filter.ref = f'A1:{get_column_letter(len(KOLOMMEN))}{ws.max_row}'

    uitleg = boek.create_sheet('Uitleg')
    uitleg.column_dimensions['A'].width = 20
    uitleg.column_dimensions['B'].width = 78
    regels = [
        ('Wat je invult', 'in de kolom Oordeel op het tabblad Fotos'),
        ('1', 'deze foto is goed'),
        ('0', 'deze foto deugt niet, de vraag krijgt er geen'),
        ('een Commons-URL', 'gebruik deze foto in plaats van het voorstel'),
        ('vrije tekst', 'een opmerking, bijvoorbeeld dat de vraag zelf rammelt'),
        ('leeg', 'nog niet bekeken; er verandert niets'),
        ('', ''),
        ('Je hoeft niet klaar', 'Wat je beoordeeld hebt komt in de volgende ronde niet terug.'),
        ('Volgorde', 'Dagpuzzels eerst, en daarbinnen de sterkste voorstellen bovenaan.'),
        ('Naamtreffer', 'ja betekent dat de naam van het gevonden artikel in de vraag staat;'),
        ('', 'die kloppen vaker dan de rest.'),
    ]
    for n, (links, rechts) in enumerate(regels, start=1):
        uitleg.cell(n, 1, links).font = Font(bold=True)
        uitleg.cell(n, 2, rechts)

    boek.save(UIT)
    print(f'{os.path.relpath(UIT, WORTEL)}: {len(te_doen)} rijen, {meegenomen} miniaturen, '
          f'{os.path.getsize(UIT) / 1024:.0f} KB')


if __name__ == '__main__':
    main()

# -*- coding: utf-8 -*-
"""Netto — verwerk het beoordeelde fotoblad.

Draaien:  python tools/verwerk_fotoblad_afkeuren.py [pad naar het xlsx]

WAT ER IN DE KOLOM "OORDEEL" MAG STAAN
  1                 deze foto is goed
  0                 deze foto deugt niet
  een Commons-URL   gebruik deze in plaats van het voorstel
  vrije tekst       een opmerking; het script beslist er niets mee maar toont hem
  leeg              nog niet bekeken, blijft zoals het is

Een eerdere versie las "alles wat ingevuld is" als afkeuren, met een apart
vakje voor hoe ver je gekomen was. Dat botste met hoe een mens zo'n blad
invult: een 1 voor goed en een 0 voor slecht. Zo hoort het ook, want dan is een
ingevulde regel altijd een beslissing en een lege regel altijd "nog niet
bekeken". Het vakje Voortgang is daarmee overbodig geworden.

Wat eruit komt:
  vragen/fotokeuze.xlsx        kolom Keuze op 1 of 0
  fotos/handmatige_fotos.json  de foto's die je met een URL hebt aangewezen;
                               die winnen het van elk automatisch voorstel
"""

import json
import os
import re
import shutil
import ssl
import sys
import time
import urllib.parse
import urllib.request
from datetime import date

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

WORTEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STANDAARD = os.path.join(WORTEL, 'vragen', 'fotos_afkeuren.xlsx')
KEUZE = os.path.join(WORTEL, 'vragen', 'fotokeuze.xlsx')
HANDMATIG = os.path.join(WORTEL, 'fotos', 'handmatige_fotos.json')
AGENT = 'Netto-fotozoeker/1.0 (penoftafel@gmail.com)'

try:
    import certifi
    CONTEXT = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    CONTEXT = ssl.create_default_context()


def bestandsnaam(url):
    """De naam van het Commons-bestand uit een geplakt adres.

    Een gewone link eindigt op de bestandsnaam. Een miniatuur heeft er een maat
    achter geplakt: .../thumb/d/d6/Naam.jpg/250px-Naam.jpg. In dat geval staat
    de echte naam een stap eerder in het pad."""
    pad = urllib.parse.urlparse(str(url)).path
    delen = [urllib.parse.unquote(d) for d in pad.split('/') if d]
    if not delen:
        return None
    if 'thumb' in delen and len(delen) >= 2:
        return delen[-2]
    return delen[-1]


def commons_gegevens(bestanden):
    uit = {}
    for start in range(0, len(bestanden), 50):
        deel = bestanden[start:start + 50]
        vraag = {'action': 'query', 'format': 'json', 'formatversion': '2',
                 'titles': '|'.join('File:' + b for b in deel), 'prop': 'imageinfo',
                 'iiprop': 'url|extmetadata|mime', 'iiurlwidth': '330'}
        url = 'https://commons.wikimedia.org/w/api.php?' + urllib.parse.urlencode(vraag)
        verzoek = urllib.request.Request(url, headers={'User-Agent': AGENT})
        try:
            with urllib.request.urlopen(verzoek, timeout=30, context=CONTEXT) as a:
                data = json.loads(a.read().decode('utf-8'))
        except Exception as fout:
            print(f'  Commons antwoordde niet: {fout}')
            continue
        for pagina in (data.get('query', {}) or {}).get('pages', []) or []:
            info = (pagina.get('imageinfo') or [{}])[0]
            if not info:
                continue
            meta = info.get('extmetadata') or {}
            def tekst(sleutel):
                blok = meta.get(sleutel) or {}
                ruw = blok.get('value', '') if isinstance(blok, dict) else ''
                return ' '.join(re.sub(r'<[^>]+>', ' ', str(ruw)).replace('&amp;', '&').split())
            uit[pagina['title']] = {
                'titel': pagina['title'],
                'pagina': info.get('descriptionurl', ''),
                'url': info.get('thumburl') or info.get('url', ''),
                'licentie': tekst('LicenseShortName') or 'onbekend',
                'maker': (tekst('Artist') or 'onbekend')[:120],
            }
        time.sleep(0.35)
    return uit


def main():
    pad = sys.argv[1] if len(sys.argv) > 1 else STANDAARD
    if not os.path.exists(pad):
        print(f'niet gevonden: {pad}')
        return
    blad = openpyxl.load_workbook(pad, read_only=True)['Fotos']
    rijen = list(blad.iter_rows(values_only=True))
    k = {naam: n for n, naam in enumerate(rijen[0])}

    goed, weg, urls, opmerkingen = {}, {}, {}, []
    for nummer, rij in enumerate(rijen[1:], start=2):
        if rij[k['Nr']] is None:
            continue
        nr = int(rij[k['Nr']])
        # Niet "of leeg" gebruiken: een cel met een 0 erin is in Python onwaar,
        # en dan verdwijnt precies de afkeuring die je bedoelde.
        rauw = rij[k['Oordeel']]
        waarde = '' if rauw is None else str(rauw).strip()
        if not waarde:
            continue
        if waarde in ('1', '1.0'):
            goed[nr] = (rij[k['Vraag']], rij[k['Bestand']])
        elif waarde in ('0', '0.0'):
            weg[nr] = rij[k['Vraag']]
        elif waarde.startswith('http'):
            urls[nr] = (waarde, rij[k['Vraag']])
        else:
            opmerkingen.append((nummer, nr, rij[k['Vraag']], waarde))

    print(f'{len(goed)} goedgekeurd | {len(weg)} weggestreept | '
          f'{len(urls)} met een eigen foto | {len(opmerkingen)} opmerkingen')

    if urls:
        print('\neigen foto\'s opzoeken op Commons ...')
        namen = {nr: bestandsnaam(u) for nr, (u, _) in urls.items()}
        gegevens = commons_gegevens(sorted({n for n in namen.values() if n}))
        handmatig = json.load(open(HANDMATIG, encoding='utf-8')) if os.path.exists(HANDMATIG) else {}
        for nr, naam in namen.items():
            info = gegevens.get('File:' + (naam or '').replace('_', ' '))
            if not info:
                print(f'  Nr {nr}: niet gevonden op Commons ({naam})')
                continue
            handmatig[str(nr)] = {'vraag': urls[nr][1], **info}
            print(f'  Nr {nr}: {info["titel"].replace("File:", "")[:52]} | {info["licentie"]}')
        with open(HANDMATIG, 'w', encoding='utf-8') as f:
            json.dump(handmatig, f, ensure_ascii=False, indent=1)
        print(f'-> {os.path.relpath(HANDMATIG, WORTEL)}')

    if opmerkingen:
        print('\nopmerkingen, die verwerkt niemand automatisch:')
        for nummer, nr, vraag, tekst in opmerkingen:
            print(f'  rij {nummer} (Nr {nr}): {tekst}')
            print(f'      {str(vraag)[:88]}')

    if not goed and not weg:
        return
    shutil.copy2(KEUZE, os.path.join(
        os.path.dirname(KEUZE),
        f'_backup_oordeel_{date.today():%Y-%m-%d}_{os.path.basename(KEUZE)}'))
    boek = openpyxl.load_workbook(KEUZE)
    keuzeblad = boek['Fotokeuze']
    kk = {naam: n for n, naam in enumerate([c.value for c in keuzeblad[1]], start=1)}
    gezet = 0
    for rij in range(2, keuzeblad.max_row + 1):
        nr = keuzeblad.cell(rij, kk['Nr']).value
        if nr is None:
            continue
        nr = int(nr)
        if nr in goed:
            keuzeblad.cell(rij, kk['Keuze']).value = 1
            gezet += 1
        elif nr in weg:
            keuzeblad.cell(rij, kk['Keuze']).value = 0
            gezet += 1
    boek.save(KEUZE)
    print(f'\n{gezet} rijen bijgewerkt in {os.path.relpath(KEUZE, WORTEL)}')

    # Het keuzeblad bevat alleen daily- en puzzelvragen. Een goedkeuring voor
    # een race- of breinkrakervraag zou daar dus stilletjes verdampen; die gaat
    # als handmatige keuze naar hetzelfde bestand als de geplakte adressen.
    aanwezig = {int(keuzeblad.cell(r, kk['Nr']).value)
                for r in range(2, keuzeblad.max_row + 1)
                if keuzeblad.cell(r, kk['Nr']).value is not None}
    verweesd = {nr: v for nr, v in goed.items() if nr not in aanwezig}
    if verweesd:
        print(f'{len(verweesd)} goedkeuringen staan niet in het keuzeblad; die gaan apart')
        namen = sorted({str(bestand) for _, bestand in verweesd.values() if bestand})
        gegevens = commons_gegevens(namen)
        handmatig = json.load(open(HANDMATIG, encoding='utf-8')) if os.path.exists(HANDMATIG) else {}
        for nr, (vraag, bestand) in verweesd.items():
            info = gegevens.get('File:' + str(bestand).replace('_', ' '))
            if info:
                handmatig[str(nr)] = {'vraag': vraag, **info}
        with open(HANDMATIG, 'w', encoding='utf-8') as f:
            json.dump(handmatig, f, ensure_ascii=False, indent=1)
    print('Draai nu: python tools/maak_fotobestand.py')


if __name__ == '__main__':
    main()
